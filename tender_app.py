#!/usr/bin/env python3
"""
Тендер Dashboard Backend — Flask + PostgreSQL
Ажиллуулах: python tender_app.py
"""

import os
import io
import logging
import psycopg2
import psycopg2.extras
import psycopg2.pool
import secrets
from decimal import Decimal
from werkzeug.security import generate_password_hash, check_password_hash
import json
from datetime import datetime
from functools import wraps
from flask import Flask, request, jsonify, session, send_file
from flask_limiter import Limiter
from flask_limiter.util import get_remote_address
import openpyxl
import anthropic
import openai
import pytesseract
from PIL import Image
from pdf2image import convert_from_bytes

logging.basicConfig(level=logging.INFO, format="%(asctime)s [%(levelname)s] %(message)s")

app = Flask(__name__, static_folder=None)
app.secret_key = os.environ.get("SECRET_KEY", secrets.token_hex(32))

limiter = Limiter(get_remote_address, app=app, default_limits=[])

DB_CONFIG = {
    "host": os.environ.get("DB_HOST", "localhost"),
    "port": os.environ.get("DB_PORT", "5432"),
    "dbname": os.environ.get("DB_NAME", "tender_db"),
    "user": os.environ.get("DB_USER", "tender_admin"),
    "password": os.environ.get("DB_PASSWORD", "admin_pass"),
}

BASE_DIR = os.path.dirname(os.path.abspath(__file__))
OVERVIEW_PATH = os.path.join(BASE_DIR, "overview.html")
DASHBOARD_PATH = os.path.join(BASE_DIR, "tender-dashboard.html")
PROCUREMENT_PATH = os.path.join(BASE_DIR, "procurement.html")
JOINED_PATH = os.path.join(BASE_DIR, "joined.html")
SUGGESTED_PATH = os.path.join(BASE_DIR, "suggested.html")

# ============================================================
# AI CLIENTS — нэг удаа үүсгэх
# ============================================================

_anthropic_key = os.environ.get("ANTHROPIC_API_KEY", "")
_openai_key = os.environ.get("OPENAI_API_KEY", "")

claude_client = anthropic.Anthropic(api_key=_anthropic_key) if _anthropic_key else None
openai_client = openai.OpenAI(api_key=_openai_key) if _openai_key else None

CLAUDE_MODEL = os.environ.get("CLAUDE_MODEL", "claude-sonnet-4-5")


# ============================================================
# DATABASE
# ============================================================

_pool: psycopg2.pool.ThreadedConnectionPool | None = None

def get_pool() -> psycopg2.pool.ThreadedConnectionPool:
    global _pool
    if _pool is None:
        _pool = psycopg2.pool.ThreadedConnectionPool(2, 20, **DB_CONFIG)
    return _pool

def get_db():
    conn = get_pool().getconn()
    conn.autocommit = False
    return conn

def release_db(conn):
    get_pool().putconn(conn)

def query(sql, params=None, fetchone=False, fetchall=False, commit=False):
    conn = get_db()
    try:
        cur = conn.cursor(cursor_factory=psycopg2.extras.RealDictCursor)
        cur.execute(sql, params)
        result = None
        if fetchone:
            result = cur.fetchone()
        elif fetchall:
            result = cur.fetchall()
        if commit:
            conn.commit()
        return result
    except Exception:
        conn.rollback()
        raise
    finally:
        release_db(conn)

def init_db():
    conn = get_db()
    cur = conn.cursor()

    cur.execute("""
        CREATE TABLE IF NOT EXISTS users (
            id SERIAL PRIMARY KEY,
            username TEXT UNIQUE NOT NULL,
            password_hash TEXT NOT NULL,
            full_name TEXT NOT NULL,
            role TEXT NOT NULL DEFAULT 'viewer',
            created_at TIMESTAMP DEFAULT NOW()
        )
    """)

    cur.execute("""
        CREATE TABLE IF NOT EXISTS tenders (
            id SERIAL PRIMARY KEY,
            tender_no TEXT UNIQUE NOT NULL,
            name TEXT,
            organization TEXT,
            type TEXT,
            method TEXT,
            deadline TEXT,
            electronic TEXT,
            link TEXT,
            raw_data TEXT,
            created_at TIMESTAMP DEFAULT NOW(),
            updated_at TIMESTAMP DEFAULT NOW(),
            created_by INTEGER REFERENCES users(id)
        )
    """)

    cur.execute("""
        CREATE TABLE IF NOT EXISTS upload_history (
            id SERIAL PRIMARY KEY,
            user_id INTEGER NOT NULL REFERENCES users(id),
            filename TEXT,
            total_in_file INTEGER DEFAULT 0,
            inserted INTEGER DEFAULT 0,
            updated INTEGER DEFAULT 0,
            unchanged INTEGER DEFAULT 0,
            uploaded_at TIMESTAMP DEFAULT NOW()
        )
    """)

    # Худалдан авалтын төлөвлөгөө
    cur.execute("""
        CREATE TABLE IF NOT EXISTS procurement_plans (
            id SERIAL PRIMARY KEY,
            row_num INTEGER,
            ts_code TEXT,
            name TEXT,
            ts_type TEXT,
            funding_source TEXT,
            budget_amount NUMERIC(18,2) DEFAULT 0,
            procurement_method TEXT,
            tender_announce_date TEXT,
            contract_award_date TEXT,
            contract_end_date TEXT,
            annual_funding NUMERIC(18,2) DEFAULT 0,
            sustainable_criteria TEXT,
            description TEXT,
            created_date TEXT,
            sub_category TEXT,
            organization TEXT,
            plan_year INTEGER,
            created_by INTEGER REFERENCES users(id),
            created_at TIMESTAMP DEFAULT NOW(),
            updated_at TIMESTAMP DEFAULT NOW()
        )
    """)

    cur.execute("""
        CREATE TABLE IF NOT EXISTS procurement_upload_history (
            id SERIAL PRIMARY KEY,
            user_id INTEGER NOT NULL REFERENCES users(id),
            filename TEXT,
            organization TEXT,
            plan_year INTEGER,
            total_rows INTEGER DEFAULT 0,
            inserted INTEGER DEFAULT 0,
            updated INTEGER DEFAULT 0,
            uploaded_at TIMESTAMP DEFAULT NOW()
        )
    """)

    # Оролцсон тендерүүд
    cur.execute("""
        CREATE TABLE IF NOT EXISTS joined_tenders (
            id SERIAL PRIMARY KEY,
            tender_no TEXT,
            code TEXT,
            name TEXT,
            organization TEXT,
            tender_type TEXT,
            method TEXT,
            status TEXT,
            deadline TEXT,
            link TEXT,
            created_by INTEGER REFERENCES users(id),
            created_at TIMESTAMP DEFAULT NOW()
        )
    """)
    cur.execute("CREATE UNIQUE INDEX IF NOT EXISTS idx_joined_tender_no ON joined_tenders(tender_no)")

    # Тендерийн нэмэлт мэдээлэл (шаардлага, тайлбар, унасан шалтгаан)
    cur.execute("""
        CREATE TABLE IF NOT EXISTS tender_details (
            id SERIAL PRIMARY KEY,
            tender_no TEXT UNIQUE NOT NULL,
            requirements TEXT,
            fail_reason TEXT,
            notes TEXT,
            ocr_text TEXT,
            updated_by INTEGER REFERENCES users(id),
            updated_at TIMESTAMP DEFAULT NOW()
        )
    """)

    # Профайл түлхүүр үгс
    cur.execute("""
        CREATE TABLE IF NOT EXISTS profile_keywords (
            id SERIAL PRIMARY KEY,
            keyword TEXT UNIQUE NOT NULL,
            weight NUMERIC(4,2) DEFAULT 1.0,
            category TEXT,
            created_at TIMESTAMP DEFAULT NOW()
        )
    """)

    # Шалгарсан тендерүүд
    cur.execute("""
        CREATE TABLE IF NOT EXISTS won_tenders (
            id SERIAL PRIMARY KEY,
            tender_no TEXT UNIQUE,
            code TEXT,
            name TEXT,
            organization TEXT,
            tender_type TEXT,
            method TEXT,
            status TEXT,
            deadline TEXT,
            link TEXT,
            created_by INTEGER REFERENCES users(id),
            created_at TIMESTAMP DEFAULT NOW()
        )
    """)

    # Vector embedding хүснэгт
    cur.execute("CREATE EXTENSION IF NOT EXISTS vector")
    cur.execute("""
        CREATE TABLE IF NOT EXISTS tender_embeddings (
            id SERIAL PRIMARY KEY,
            tender_no TEXT NOT NULL,
            chunk_index INTEGER DEFAULT 0,
            chunk_text TEXT,
            embedding vector(1536),
            source TEXT DEFAULT 'ocr',
            created_at TIMESTAMP DEFAULT NOW()
        )
    """)
    cur.execute("CREATE INDEX IF NOT EXISTS idx_emb_tender ON tender_embeddings(tender_no)")

    # Index-ууд
    cur.execute("CREATE INDEX IF NOT EXISTS idx_tenders_tender_no ON tenders(tender_no)")
    cur.execute("CREATE INDEX IF NOT EXISTS idx_tenders_deadline ON tenders(deadline)")
    cur.execute("CREATE INDEX IF NOT EXISTS idx_tenders_organization ON tenders(organization)")
    cur.execute("CREATE INDEX IF NOT EXISTS idx_tenders_type ON tenders(type)")
    cur.execute("CREATE INDEX IF NOT EXISTS idx_procurement_ts_code ON procurement_plans(ts_code)")
    cur.execute("CREATE INDEX IF NOT EXISTS idx_procurement_org ON procurement_plans(organization)")
    cur.execute("CREATE INDEX IF NOT EXISTS idx_procurement_type ON procurement_plans(ts_type)")
    cur.execute("CREATE INDEX IF NOT EXISTS idx_procurement_year ON procurement_plans(plan_year)")
    cur.execute("CREATE INDEX IF NOT EXISTS idx_tenders_name_lower ON tenders(LOWER(TRIM(name)))")
    cur.execute("CREATE INDEX IF NOT EXISTS idx_procurement_name_lower ON procurement_plans(LOWER(TRIM(name)))")

    conn.commit()

    # Анхдагч admin
    cur.execute("SELECT id FROM users WHERE username = %s", ("admin",))
    if not cur.fetchone():
        pw_hash = generate_password_hash("admin123")
        cur.execute(
            "INSERT INTO users (username, password_hash, full_name, role) VALUES (%s, %s, %s, %s)",
            ("admin", pw_hash, "Админ", "admin")
        )
        conn.commit()
        print("[tender] Анхдагч хэрэглэгч: admin / admin123")

    release_db(conn)


def dec(obj):
    """JSON serialize helper — Decimal болон datetime хөрвүүлэх"""
    if isinstance(obj, Decimal):
        return float(obj)
    if isinstance(obj, datetime):
        return obj.isoformat()
    return obj


# ============================================================
# MIDDLEWARE
# ============================================================

def login_required(f):
    @wraps(f)
    def decorated(*args, **kwargs):
        if "user_id" not in session:
            return jsonify({"error": "Нэвтрэх шаардлагатай"}), 401
        return f(*args, **kwargs)
    return decorated

def editor_required(f):
    """Зөвхөн admin, editor эрхтэй хэрэглэгч"""
    @wraps(f)
    def decorated(*args, **kwargs):
        if "user_id" not in session:
            return jsonify({"error": "Нэвтрэх шаардлагатай"}), 401
        if session.get("role") not in ("admin", "editor"):
            return jsonify({"error": "Мэдээлэл оруулах эрхгүй. admin эсвэл editor эрх шаардлагатай."}), 403
        return f(*args, **kwargs)
    return decorated


# ============================================================
# AUTH API
# ============================================================

@app.route("/api/login", methods=["POST"])
@limiter.limit("10 per minute")
def login():
    data = request.get_json()
    username = data.get("username", "").strip()
    password = data.get("password", "").strip()

    if not username or not password:
        return jsonify({"error": "Хэрэглэгчийн нэр, нууц үг оруулна уу"}), 400

    user = query(
        "SELECT id, username, full_name, role, password_hash FROM users WHERE username = %s",
        (username,), fetchone=True
    )

    if not user or not check_password_hash(user["password_hash"], password):
        return jsonify({"error": "Хэрэглэгчийн нэр эсвэл нууц үг буруу"}), 401

    session["user_id"] = user["id"]
    session["username"] = user["username"]
    session["full_name"] = user["full_name"]
    session["role"] = user["role"]

    return jsonify({
        "success": True,
        "user": {
            "username": user["username"],
            "full_name": user["full_name"],
            "role": user["role"]
        }
    })

@app.route("/api/logout", methods=["POST"])
def logout():
    session.clear()
    return jsonify({"success": True})

@app.route("/api/me")
def me():
    if "user_id" not in session:
        return jsonify({"logged_in": False}), 401
    return jsonify({
        "logged_in": True,
        "user": {
            "username": session["username"],
            "full_name": session["full_name"],
            "role": session["role"]
        }
    })


# ============================================================
# ХЭРЭГЛЭГЧ УДИРДЛАГА (admin only)
# ============================================================

@app.route("/api/users", methods=["GET"])
@login_required
def list_users():
    if session.get("role") != "admin":
        return jsonify({"error": "Зөвхөн админ"}), 403
    users = query("SELECT id, username, full_name, role, created_at FROM users", fetchall=True)
    return jsonify(users)

@app.route("/api/users", methods=["POST"])
@login_required
def create_user():
    if session.get("role") != "admin":
        return jsonify({"error": "Зөвхөн админ"}), 403

    data = request.get_json()
    username = data.get("username", "").strip()
    password = data.get("password", "").strip()
    full_name = data.get("full_name", "").strip()
    role = data.get("role", "viewer")

    if role not in ("admin", "editor", "viewer"):
        return jsonify({"error": "Эрх буруу. admin, editor, viewer сонгоно уу"}), 400
    if not username or not password or not full_name:
        return jsonify({"error": "Бүх талбарыг бөглөнө үү"}), 400

    pw_hash = generate_password_hash(password)
    try:
        query(
            "INSERT INTO users (username, password_hash, full_name, role) VALUES (%s, %s, %s, %s)",
            (username, pw_hash, full_name, role), commit=True
        )
    except psycopg2.errors.UniqueViolation:
        return jsonify({"error": "Хэрэглэгчийн нэр давхцаж байна"}), 409
    return jsonify({"success": True}), 201

@app.route("/api/users/<int:user_id>", methods=["DELETE"])
@login_required
def delete_user(user_id):
    if session.get("role") != "admin":
        return jsonify({"error": "Зөвхөн админ"}), 403
    if user_id == session.get("user_id"):
        return jsonify({"error": "Өөрийгөө устгах боломжгүй"}), 400
    query("DELETE FROM users WHERE id = %s", (user_id,), commit=True)
    return jsonify({"success": True})


# ============================================================
# ТЕНДЕР UPLOAD — diff/merge логик
# ============================================================

TENDER_FIELDS = ("name", "organization", "type", "method", "deadline", "electronic", "link")

@app.route("/api/tenders/upload", methods=["POST"])
@editor_required
def upload_tenders():
    data = request.get_json()
    tenders = data.get("tenders", [])
    filename = data.get("filename", "unknown.json")

    if not tenders or not isinstance(tenders, list):
        return jsonify({"error": "Тендерийн мэдээлэл хоосон байна"}), 400

    conn = get_db()
    cur = conn.cursor(cursor_factory=psycopg2.extras.RealDictCursor)
    user_id = session["user_id"]
    now = datetime.utcnow().isoformat()

    inserted = 0
    updated = 0
    unchanged = 0
    changes = []

    try:
        for t in tenders:
            tender_no = t.get("tender_no", "").strip()
            if not tender_no:
                continue

            cur.execute(
                "SELECT id, name, organization, type, method, deadline, electronic, link FROM tenders WHERE tender_no = %s",
                (tender_no,)
            )
            existing = cur.fetchone()

            if existing is None:
                cur.execute("""
                    INSERT INTO tenders (tender_no, name, organization, type, method, deadline, electronic, link, raw_data, created_by, created_at, updated_at)
                    VALUES (%s, %s, %s, %s, %s, %s, %s, %s, %s, %s, %s, %s)
                """, (
                    tender_no,
                    t.get("name"), t.get("organization"), t.get("type"),
                    t.get("method"), t.get("deadline"), t.get("electronic"), t.get("link"),
                    json.dumps(t, ensure_ascii=False),
                    user_id, now, now
                ))
                inserted += 1
                changes.append({"tender_no": tender_no, "action": "new", "name": t.get("name", "")})
            else:
                diffs = {}
                for field in TENDER_FIELDS:
                    old_val = existing[field] or ""
                    new_val = t.get(field, "") or ""
                    if str(old_val) != str(new_val):
                        diffs[field] = {"old": old_val, "new": new_val}

                if diffs:
                    cur.execute("""
                        UPDATE tenders
                        SET name=%s, organization=%s, type=%s, method=%s, deadline=%s, electronic=%s, link=%s, raw_data=%s, updated_at=%s
                        WHERE tender_no=%s
                    """, (
                        t.get("name"), t.get("organization"), t.get("type"),
                        t.get("method"), t.get("deadline"), t.get("electronic"), t.get("link"),
                        json.dumps(t, ensure_ascii=False),
                        now, tender_no
                    ))
                    updated += 1
                    changes.append({"tender_no": tender_no, "action": "updated", "name": t.get("name", ""), "diffs": diffs})
                else:
                    unchanged += 1

        cur.execute("""
            INSERT INTO upload_history (user_id, filename, total_in_file, inserted, updated, unchanged)
            VALUES (%s, %s, %s, %s, %s, %s)
        """, (user_id, filename, len(tenders), inserted, updated, unchanged))

        conn.commit()
    except Exception:
        conn.rollback()
        raise
    finally:
        release_db(conn)

    return jsonify({
        "success": True,
        "summary": {
            "total_in_file": len(tenders),
            "inserted": inserted,
            "updated": updated,
            "unchanged": unchanged
        },
        "changes": changes[:50]
    })


# ============================================================
# ТЕНДЕР QUERY API
# ============================================================

@app.route("/api/tenders")
@login_required
def get_tenders():
    total = query("SELECT COUNT(*) as cnt FROM tenders", fetchone=True)["cnt"]

    limit = min(int(request.args.get("limit", 5000)), 10000)
    offset = int(request.args.get("offset", 0))

    rows = query("""
        SELECT tender_no, name, organization, type, method, deadline, electronic, link,
               created_at, updated_at
        FROM tenders ORDER BY deadline ASC LIMIT %s OFFSET %s
    """, (limit, offset), fetchall=True)

    return jsonify({
        "total": total,
        "tenders": rows
    })

@app.route("/api/procurement/pending")
@login_required
def procurement_pending():
    """Төлөвлөгөөнөөс тендер зарлагдаагүй эсвэл хугацаа дуусаагүй.
    Тулгах: 1) ТШ код = тендер дугаарын эхний 2 хэсэг
            2) Нэр бүрэн тохирох
    Оновчилсон: нэг удаа тулгаж temp table-д хадгална.
    """
    search = request.args.get("search", "")
    ts_type = request.args.get("type", "")
    org = request.args.get("organization", "")
    year = request.args.get("year", "")
    status = request.args.get("status", "")
    limit = min(int(request.args.get("limit", 100)), 5000)
    offset = int(request.args.get("offset", 0))

    conn = get_db()
    cur = conn.cursor(cursor_factory=psycopg2.extras.RealDictCursor)

    try:
        # 1) Тулгалтын хүснэгт — нэг удаа бүтээх
        cur.execute("""
            CREATE TEMP TABLE matched_plans ON COMMIT DROP AS
            SELECT p.*,
                   t.tender_no as matched_tender,
                   t.deadline as tender_deadline,
                   t.method as tender_method,
                   t.link as tender_link,
                   CASE WHEN t.tender_no IS NULL THEN 'not_announced' ELSE 'active' END as tender_status
            FROM procurement_plans p
            LEFT JOIN (
                SELECT DISTINCT ON (code)
                    SPLIT_PART(tender_no, '/', 1) || '/' || SPLIT_PART(tender_no, '/', 2) as code,
                    tender_no, deadline, method, link, name
                FROM tenders
                ORDER BY code, deadline DESC
            ) t ON (p.ts_code = t.code OR LOWER(TRIM(p.name)) = LOWER(TRIM(t.name)))
            WHERE (t.tender_no IS NULL OR t.deadline::timestamp >= NOW())
        """)

        # 2) Шүүлт
        where = ["1=1"]
        params = []
        if search:
            where.append("(name ILIKE %s OR ts_code ILIKE %s)")
            params.extend([f"%{search}%"] * 2)
        if ts_type:
            where.append("ts_type = %s")
            params.append(ts_type)
        if org:
            where.append("organization = %s")
            params.append(org)
        if year:
            where.append("plan_year = %s")
            params.append(int(year))
        if status:
            where.append("tender_status = %s")
            params.append(status)

        where_sql = " AND ".join(where)

        # 3) Stats — нэг query
        cur.execute(f"""
            SELECT COUNT(*) as total,
                   COUNT(CASE WHEN tender_status = 'not_announced' THEN 1 END) as not_announced,
                   COUNT(CASE WHEN tender_status = 'active' THEN 1 END) as announced
            FROM matched_plans WHERE {where_sql}
        """, params)
        stats = cur.fetchone()

        # 4) Data
        cur.execute(f"""
            SELECT * FROM matched_plans
            WHERE {where_sql}
            ORDER BY row_num ASC LIMIT %s OFFSET %s
        """, params + [limit, offset])
        rows = cur.fetchall()

        from decimal import Decimal
        for r in rows:
            for k, v in r.items():
                if isinstance(v, Decimal):
                    r[k] = float(v)

        conn.commit()
    except Exception:
        conn.rollback()
        raise
    finally:
        release_db(conn)

    return jsonify({
        "total": stats["total"],
        "plans": rows,
        "stats": {
            "total": stats["total"],
            "not_announced": stats["not_announced"],
            "announced": stats["announced"]
        }
    })


@app.route("/api/tenders/stats")
@login_required
def tender_stats():
    total = query("SELECT COUNT(*) as cnt FROM tenders", fetchone=True)["cnt"]
    orgs = query("SELECT COUNT(DISTINCT organization) as cnt FROM tenders", fetchone=True)["cnt"]
    types = query("SELECT type, COUNT(*) as cnt FROM tenders GROUP BY type", fetchall=True)
    methods = query("SELECT method, COUNT(*) as cnt FROM tenders GROUP BY method", fetchall=True)
    return jsonify({
        "total": total,
        "organizations": orgs,
        "types": types,
        "methods": methods
    })

@app.route("/api/tenders/history")
@login_required
def upload_history():
    rows = query("""
        SELECT h.*, u.full_name, u.username
        FROM upload_history h JOIN users u ON h.user_id = u.id
        ORDER BY h.uploaded_at DESC LIMIT 50
    """, fetchall=True)
    return jsonify(rows)


# ============================================================
# ХУДАЛДАН АВАЛТЫН ТӨЛӨВЛӨГӨӨ
# ============================================================

@app.route("/api/procurement/upload", methods=["POST"])
@editor_required
def upload_procurement():
    if "file" not in request.files:
        return jsonify({"error": "Excel файл сонгоно уу"}), 400

    file = request.files["file"]
    if not file.filename.lower().endswith((".xlsx", ".xls")):
        return jsonify({"error": "Зөвхөн .xlsx файл дэмжигдэнэ"}), 400

    organization = request.form.get("organization", "").strip()
    plan_year = int(request.form.get("plan_year", datetime.now().year))

    if not organization:
        return jsonify({"error": "Байгууллагын нэр оруулна уу"}), 400

    try:
        wb = openpyxl.load_workbook(io.BytesIO(file.read()), read_only=True, data_only=True)

        # sheet1 эсвэл эхний sheet-ээс уншина
        ws = wb["sheet1"] if "sheet1" in wb.sheetnames else wb[wb.sheetnames[0]]

        # Header мөр тодорхойлох
        headers = []
        for cell in next(ws.iter_rows(min_row=1, max_row=1)):
            headers.append(str(cell.value or "").strip())

        # Баганы индекс олох
        col_map = {}
        expected = {
            "№": "row_num",
            "ТШ код": "ts_code",
            "нэр төрөл": "name",
            "ТШ төрөл": "ts_type",
            "эх үүсвэр": "funding_source",
            "Төсөвт өртөг": "budget_amount",
            "журам": "procurement_method",
            "Тендер зарлах": "tender_announce_date",
            "эрх олгох": "contract_award_date",
            "дуусгавар": "contract_end_date",
            "санхүүжих дүн": "annual_funding",
            "Тогтвортой": "sustainable_criteria",
            "Тайлбар": "description",
            "Үүсгэсэн огноо": "created_date",
        }
        for i, h in enumerate(headers):
            for key, field in expected.items():
                if key.lower() in h.lower():
                    col_map[field] = i
                    break

        conn = get_db()
        cur = conn.cursor()
        user_id = session["user_id"]
        now = datetime.utcnow().isoformat()

        inserted = 0
        updated = 0
        total = 0

        # Дэд ангилал-тай sheet байвал уншина
        sub_cats = {}
        for sname in ["Бараа", "Ажил", "Үйлчилгээ"]:
            if sname in wb.sheetnames:
                sub_ws = wb[sname]
                sub_headers = [str(c.value or "").strip() for c in next(sub_ws.iter_rows(min_row=1, max_row=1))]
                code_idx = next((i for i, h in enumerate(sub_headers) if "код" in h.lower()), None)
                sub_idx = next((i for i, h in enumerate(sub_headers) if "ангилал" in h.lower()), None)
                if code_idx is not None and sub_idx is not None:
                    for row in sub_ws.iter_rows(min_row=2, values_only=True):
                        code = str(row[code_idx] or "").strip()
                        sub = str(row[sub_idx] or "").strip()
                        if code and sub:
                            sub_cats[code] = sub

        for row in ws.iter_rows(min_row=2, values_only=True):
            vals = list(row)
            if not vals or len(vals) < 3:
                continue

            ts_code = str(vals[col_map.get("ts_code", 1)] or "").strip()
            if not ts_code:
                continue

            total += 1
            name = str(vals[col_map.get("name", 2)] or "").strip()
            ts_type = str(vals[col_map.get("ts_type", 3)] or "").strip()
            funding = str(vals[col_map.get("funding_source", 4)] or "").strip()
            budget = vals[col_map.get("budget_amount", 5)] or 0
            method = str(vals[col_map.get("procurement_method", 6)] or "").strip()
            t_announce = str(vals[col_map.get("tender_announce_date", 7)] or "").strip()
            c_award = str(vals[col_map.get("contract_award_date", 8)] or "").strip()
            c_end = str(vals[col_map.get("contract_end_date", 9)] or "").strip()
            ann_fund = vals[col_map.get("annual_funding", 10)] or 0
            sustain = str(vals[col_map.get("sustainable_criteria", 11)] or "").strip()
            desc = str(vals[col_map.get("description", 12)] or "").strip()
            c_date = str(vals[col_map.get("created_date", 13)] or "").strip()
            row_num_val = vals[col_map.get("row_num", 0)] or 0

            try:
                budget = float(budget)
            except (ValueError, TypeError):
                budget = 0
            try:
                ann_fund = float(ann_fund)
            except (ValueError, TypeError):
                ann_fund = 0
            try:
                row_num_val = int(row_num_val)
            except (ValueError, TypeError):
                row_num_val = 0

            sub_cat = sub_cats.get(ts_code, "")

            # Upsert
            cur.execute(
                "SELECT id FROM procurement_plans WHERE ts_code = %s AND organization = %s AND plan_year = %s",
                (ts_code, organization, plan_year)
            )
            existing = cur.fetchone()

            if existing:
                cur.execute("""
                    UPDATE procurement_plans SET
                        row_num=%s, name=%s, ts_type=%s, funding_source=%s, budget_amount=%s,
                        procurement_method=%s, tender_announce_date=%s, contract_award_date=%s,
                        contract_end_date=%s, annual_funding=%s, sustainable_criteria=%s,
                        description=%s, created_date=%s, sub_category=%s, updated_at=%s
                    WHERE ts_code=%s AND organization=%s AND plan_year=%s
                """, (
                    row_num_val, name, ts_type, funding, budget,
                    method, t_announce, c_award, c_end, ann_fund,
                    sustain, desc, c_date, sub_cat, now,
                    ts_code, organization, plan_year
                ))
                updated += 1
            else:
                cur.execute("""
                    INSERT INTO procurement_plans
                        (row_num, ts_code, name, ts_type, funding_source, budget_amount,
                         procurement_method, tender_announce_date, contract_award_date,
                         contract_end_date, annual_funding, sustainable_criteria,
                         description, created_date, sub_category, organization,
                         plan_year, created_by, created_at, updated_at)
                    VALUES (%s,%s,%s,%s,%s,%s,%s,%s,%s,%s,%s,%s,%s,%s,%s,%s,%s,%s,%s,%s)
                """, (
                    row_num_val, ts_code, name, ts_type, funding, budget,
                    method, t_announce, c_award, c_end, ann_fund,
                    sustain, desc, c_date, sub_cat, organization,
                    plan_year, user_id, now, now
                ))
                inserted += 1

        # Upload history
        cur.execute("""
            INSERT INTO procurement_upload_history
                (user_id, filename, organization, plan_year, total_rows, inserted, updated)
            VALUES (%s, %s, %s, %s, %s, %s, %s)
        """, (user_id, file.filename, organization, plan_year, total, inserted, updated))

        conn.commit()
        release_db(conn)
        wb.close()

        return jsonify({
            "success": True,
            "summary": {
                "total": total,
                "inserted": inserted,
                "updated": updated,
                "organization": organization,
                "plan_year": plan_year
            }
        })
    except Exception as e:
        return jsonify({"error": f"Файл боловсруулж чадсангүй: {str(e)}"}), 500


@app.route("/api/procurement")
@login_required
def get_procurement():
    org = request.args.get("organization", "")
    year = request.args.get("year", "")
    ts_type = request.args.get("type", "")
    search = request.args.get("search", "")
    limit = min(int(request.args.get("limit", 5000)), 10000)
    offset = int(request.args.get("offset", 0))

    where = []
    params = []
    if org:
        where.append("organization = %s")
        params.append(org)
    if year:
        where.append("plan_year = %s")
        params.append(int(year))
    if ts_type:
        where.append("ts_type = %s")
        params.append(ts_type)
    if search:
        where.append("(name ILIKE %s OR ts_code ILIKE %s OR sub_category ILIKE %s)")
        params.extend([f"%{search}%"] * 3)

    where_sql = " AND ".join(where) if where else "1=1"

    total = query(f"SELECT COUNT(*) as cnt FROM procurement_plans WHERE {where_sql}",
                  params, fetchone=True)["cnt"]

    rows = query(f"""
        SELECT * FROM procurement_plans WHERE {where_sql}
        ORDER BY row_num ASC LIMIT %s OFFSET %s
    """, params + [limit, offset], fetchall=True)

    # Numeric -> float (JSON serialization)
    from decimal import Decimal
    for r in rows:
        for k, v in r.items():
            if isinstance(v, Decimal):
                r[k] = float(v)

    return jsonify({"total": total, "plans": rows})


@app.route("/api/procurement/stats")
@login_required
def procurement_stats():
    orgs = query("SELECT DISTINCT organization FROM procurement_plans ORDER BY organization",
                 fetchall=True)
    years = query("SELECT DISTINCT plan_year FROM procurement_plans ORDER BY plan_year DESC",
                  fetchall=True)
    types = query("""
        SELECT ts_type, COUNT(*) as cnt, COALESCE(SUM(budget_amount),0) as total_budget
        FROM procurement_plans GROUP BY ts_type ORDER BY cnt DESC
    """, fetchall=True)
    total = query("SELECT COUNT(*) as cnt, COALESCE(SUM(budget_amount),0) as total_budget FROM procurement_plans",
                  fetchone=True)

    from decimal import Decimal
    for t in types:
        for k, v in t.items():
            if isinstance(v, Decimal):
                t[k] = float(v)
    for k, v in total.items():
        if isinstance(v, Decimal):
            total[k] = float(v)

    return jsonify({
        "organizations": [r["organization"] for r in orgs],
        "years": [r["plan_year"] for r in years],
        "types": types,
        "total_count": total["cnt"],
        "total_budget": total["total_budget"]
    })


@app.route("/api/procurement/history")
@login_required
def procurement_history():
    rows = query("""
        SELECT h.*, u.full_name FROM procurement_upload_history h
        JOIN users u ON h.user_id = u.id
        ORDER BY h.uploaded_at DESC LIMIT 50
    """, fetchall=True)
    return jsonify(rows)


# ============================================================
# ОРОЛЦСОН ТЕНДЕРҮҮД
# ============================================================

@app.route("/api/joined-tenders/upload", methods=["POST"])
@editor_required
def upload_joined_tenders():
    data = request.get_json()
    tenders = data.get("tenders", [])
    if not tenders or not isinstance(tenders, list):
        return jsonify({"error": "Тендерийн мэдээлэл хоосон"}), 400

    conn = get_db()
    cur = conn.cursor()
    user_id = session["user_id"]
    inserted = 0
    updated = 0

    try:
        for t in tenders:
            tender_no = (t.get("tender_no") or "").strip()
            if not tender_no:
                continue

            cur.execute("SELECT id FROM joined_tenders WHERE tender_no = %s", (tender_no,))
            existing = cur.fetchone()

            name = (t.get("name") or "").strip()
            org = (t.get("organization") or "").strip()
            code = (t.get("code") or "").strip()
            tender_type = (t.get("electronic") or t.get("type") or "").strip()
            method_val = (t.get("type") or t.get("method") or "").strip()
            status_val = (t.get("method") or "").strip()
            deadline = (t.get("deadline") or "").strip()
            if deadline == "null":
                deadline = ""
            link = (t.get("link") or "").strip()

            if existing:
                cur.execute("""
                    UPDATE joined_tenders SET code=%s, name=%s, organization=%s,
                    tender_type=%s, method=%s, status=%s, deadline=%s, link=%s
                    WHERE tender_no=%s
                """, (code, name, org, tender_type, method_val, status_val, deadline, link, tender_no))
                updated += 1
            else:
                cur.execute("""
                    INSERT INTO joined_tenders (tender_no, code, name, organization,
                    tender_type, method, status, deadline, link, created_by)
                    VALUES (%s,%s,%s,%s,%s,%s,%s,%s,%s,%s)
                """, (tender_no, code, name, org, tender_type, method_val, status_val, deadline, link, user_id))
                inserted += 1

        conn.commit()
    except Exception:
        conn.rollback()
        raise
    finally:
        release_db(conn)

    return jsonify({
        "success": True,
        "summary": {"total": len(tenders), "inserted": inserted, "updated": updated}
    })


@app.route("/api/won-tenders/upload", methods=["POST"])
@editor_required
def upload_won_tenders():
    data = request.get_json()
    tenders = data.get("tenders", [])
    if not tenders or not isinstance(tenders, list):
        return jsonify({"error": "Тендерийн мэдээлэл хоосон"}), 400

    conn = get_db()
    cur = conn.cursor()
    user_id = session["user_id"]
    inserted = 0
    updated = 0

    try:
        for t in tenders:
            tender_no = (t.get("tender_no") or "").strip()
            if not tender_no:
                continue
            code = (t.get("code") or "").strip()
            name = (t.get("name") or "").strip()
            org = (t.get("organization") or "").strip()
            tender_type = (t.get("electronic") or t.get("type") or "").strip()
            method_val = (t.get("type") or t.get("method") or "").strip()
            status_val = (t.get("method") or "").strip()
            deadline = (t.get("deadline") or "").strip()
            if deadline == "null":
                deadline = ""
            link = (t.get("link") or "").strip()

            cur.execute("SELECT id FROM won_tenders WHERE tender_no = %s", (tender_no,))
            if cur.fetchone():
                cur.execute("""
                    UPDATE won_tenders SET code=%s, name=%s, organization=%s,
                    tender_type=%s, method=%s, status=%s, deadline=%s, link=%s
                    WHERE tender_no=%s
                """, (code, name, org, tender_type, method_val, status_val, deadline, link, tender_no))
                updated += 1
            else:
                cur.execute("""
                    INSERT INTO won_tenders (tender_no, code, name, organization,
                    tender_type, method, status, deadline, link, created_by)
                    VALUES (%s,%s,%s,%s,%s,%s,%s,%s,%s,%s)
                """, (tender_no, code, name, org, tender_type, method_val, status_val, deadline, link, user_id))
                inserted += 1
        conn.commit()
    except Exception:
        conn.rollback()
        raise
    finally:
        release_db(conn)

    return jsonify({"success": True, "summary": {"total": len(tenders), "inserted": inserted, "updated": updated}})


@app.route("/api/joined-tenders")
@login_required
def get_joined_tenders():
    search = request.args.get("search", "")
    status = request.args.get("status", "")
    t_type = request.args.get("type", "")
    won = request.args.get("won", "")  # "yes", "no", ""
    limit = min(int(request.args.get("limit", 100)), 5000)
    offset = int(request.args.get("offset", 0))

    where = ["1=1"]
    params = []
    if search:
        where.append("(j.name ILIKE %s OR j.tender_no ILIKE %s OR j.organization ILIKE %s)")
        params.extend([f"%{search}%"] * 3)
    if status:
        where.append("j.status = %s")
        params.append(status)
    if t_type:
        where.append("j.tender_type = %s")
        params.append(t_type)
    if won == "yes":
        where.append("w.tender_no IS NOT NULL")
    elif won == "no":
        where.append("w.tender_no IS NULL")

    where_sql = " AND ".join(where)

    total = query(f"""
        SELECT COUNT(*) as cnt FROM joined_tenders j
        LEFT JOIN won_tenders w ON j.tender_no = w.tender_no
        WHERE {where_sql}
    """, params, fetchone=True)["cnt"]

    rows = query(f"""
        SELECT j.*,
               CASE WHEN w.tender_no IS NOT NULL THEN true ELSE false END as is_won,
               CASE WHEN d.tender_no IS NOT NULL THEN true ELSE false END as has_details,
               d.fail_reason, d.notes
        FROM joined_tenders j
        LEFT JOIN won_tenders w ON j.tender_no = w.tender_no
        LEFT JOIN tender_details d ON j.tender_no = d.tender_no
        WHERE {where_sql}
        ORDER BY j.id DESC LIMIT %s OFFSET %s
    """, params + [limit, offset], fetchall=True)

    return jsonify({"total": total, "tenders": rows})


@app.route("/api/joined-tenders/stats")
@login_required
def joined_tender_stats():
    stats = query("""
        SELECT COUNT(*) as total,
               COUNT(DISTINCT j.organization) as orgs,
               COUNT(CASE WHEN j.status = 'Үр дүн гарсан' THEN 1 END) as completed,
               COUNT(CASE WHEN j.status = 'Нээгдсэн' THEN 1 END) as opened,
               COUNT(CASE WHEN w.tender_no IS NOT NULL THEN 1 END) as won,
               COUNT(CASE WHEN j.status = 'Үр дүн гарсан' AND w.tender_no IS NULL THEN 1 END) as lost
        FROM joined_tenders j
        LEFT JOIN won_tenders w ON j.tender_no = w.tender_no
    """, fetchone=True)
    won_count = query("SELECT COUNT(*) as cnt FROM won_tenders", fetchone=True)["cnt"]
    return jsonify({"stats": stats, "won_total": won_count})


# ============================================================
# ТЕНДЕРИЙН НЭМЭЛТ МЭДЭЭЛЭЛ
# ============================================================

@app.route("/api/tender-details/<path:tender_no>")
@login_required
def get_tender_detail(tender_no):
    row = query("SELECT * FROM tender_details WHERE tender_no = %s", (tender_no,), fetchone=True)
    return jsonify(row or {"tender_no": tender_no, "requirements": "", "fail_reason": "", "notes": "", "ocr_text": ""})


@app.route("/api/tender-details/<path:tender_no>", methods=["PUT"])
@login_required
def save_tender_detail(tender_no):
    data = request.get_json()
    requirements = data.get("requirements", "")
    fail_reason = data.get("fail_reason", "")
    notes = data.get("notes", "")
    ocr_text = data.get("ocr_text", "")
    user_id = session["user_id"]

    conn = get_db()
    cur = conn.cursor()
    cur.execute("SELECT id FROM tender_details WHERE tender_no = %s", (tender_no,))
    if cur.fetchone():
        cur.execute("""UPDATE tender_details SET requirements=%s, fail_reason=%s, notes=%s, ocr_text=%s, updated_by=%s, updated_at=NOW()
                       WHERE tender_no=%s""",
                    (requirements, fail_reason, notes, ocr_text, user_id, tender_no))
    else:
        cur.execute("""INSERT INTO tender_details (tender_no, requirements, fail_reason, notes, ocr_text, updated_by)
                       VALUES (%s,%s,%s,%s,%s,%s)""",
                    (tender_no, requirements, fail_reason, notes, ocr_text, user_id))
    conn.commit()
    release_db(conn)
    return jsonify({"success": True})


# ============================================================
# OCR — PDF/Зураг уншиж шаардлага задлах
# ============================================================

@app.route("/api/ocr/upload", methods=["POST"])
@login_required
def ocr_upload():
    """PDF/зураг upload → OCR → Claude засварлаж шаардлага задлах"""
    if "file" not in request.files:
        return jsonify({"error": "Файл сонгоно уу"}), 400

    file = request.files["file"]
    tender_no = request.form.get("tender_no", "")
    fname = file.filename.lower()

    # 1) Файлаас текст гаргах
    raw_text = ""
    try:
        file_bytes = file.read()

        if fname.endswith(".pdf"):
            images = convert_from_bytes(file_bytes, dpi=300)
            for img in images:
                raw_text += pytesseract.image_to_string(img, lang="mon+eng") + "\n"
        elif fname.endswith((".png", ".jpg", ".jpeg", ".bmp", ".tiff")):
            img = Image.open(io.BytesIO(file_bytes))
            raw_text = pytesseract.image_to_string(img, lang="mon+eng")
        else:
            return jsonify({"error": "PDF эсвэл зураг файл дэмжигдэнэ"}), 400
    except Exception as e:
        return jsonify({"error": f"OCR алдаа: {str(e)}"}), 500

    if not raw_text.strip():
        return jsonify({"error": "Текст уншиж чадсангүй. Зургийн чанарыг шалгана уу."}), 400

    # 2) Claude-аар засварлаж шаардлага задлах
    client = claude_client
    if not client:
        return jsonify({
            "ocr_text": raw_text,
            "corrected": raw_text,
            "requirements": [],
            "warning": "ANTHROPIC_API_KEY тохируулаагүй — OCR текст засварлагдаагүй"
        })

    try:
        msg = client.messages.create(
            model=CLAUDE_MODEL,
            max_tokens=3000,
            messages=[{
                "role": "user",
                "content": f"""Дараах текст нь тендерийн баримт бичгээс OCR-ээр уншсан. Монгол кирилл алдааг засаж, шаардлагуудыг задал.

OCR ТЕКСТ:
{raw_text[:4000]}

ДААЛГАВАР:
1. OCR алдааг засварла (үсэг солигдсон, тэмдэгт алдаа)
2. Шаардлагуудыг жагсаалтаар гарга
3. JSON хариул:

{{"corrected_text": "засварласан бүрэн текст", "requirements": [{{"id": "3.1", "text": "шаардлагын текст", "category": "ерөнхий/техникийн/санхүүгийн/туршлагын"}}], "summary": "баримт бичгийн товч тайлбар"}}

Зөвхөн JSON хариул."""
            }]
        )

        raw = msg.content[0].text.strip()
        result = None
        try:
            result = json.loads(raw)
        except (json.JSONDecodeError, ValueError):
            start = raw.find('{')
            end = raw.rfind('}')
            if start >= 0 and end > start:
                try:
                    result = json.loads(raw[start:end+1])
                except (json.JSONDecodeError, ValueError):
                    pass

        if not result:
            result = {"corrected_text": raw_text, "requirements": [], "summary": ""}

        result["ocr_text"] = raw_text
        result["tender_no"] = tender_no

        # Автомат embedding хийх
        if tender_no:
            try:
                emb_text = result.get("corrected_text") or raw_text
                if emb_text and len(emb_text.strip()) > 10:
                    chunks = chunk_text(emb_text)
                    embeddings = embed_texts(chunks)
                    if embeddings:
                        conn2 = get_db()
                        cur2 = conn2.cursor()
                        cur2.execute("DELETE FROM tender_embeddings WHERE tender_no = %s", (tender_no,))
                        for i, (ch, em) in enumerate(zip(chunks, embeddings)):
                            cur2.execute("INSERT INTO tender_embeddings (tender_no, chunk_index, chunk_text, embedding) VALUES (%s,%s,%s,%s)",
                                         (tender_no, i, ch, str(em)))
                        conn2.commit()
                        release_db(conn2)
                        result["embedding_chunks"] = len(chunks)
            except Exception as e:
                logging.warning("OCR embedding алдаа: %s", e)

        return jsonify(result)

    except Exception as e:
        return jsonify({
            "ocr_text": raw_text,
            "corrected_text": raw_text,
            "requirements": [],
            "error": f"Claude алдаа: {str(e)}"
        })


# ============================================================
# EMBEDDING — OpenAI text-embedding-3-small + pgvector
# ============================================================



def chunk_text(text, max_len=500):
    """Текстийг chunk-ууд болгож хуваах"""
    words = text.split()
    chunks = []
    current = []
    length = 0
    for w in words:
        current.append(w)
        length += len(w) + 1
        if length >= max_len:
            chunks.append(' '.join(current))
            current = []
            length = 0
    if current:
        chunks.append(' '.join(current))
    return chunks if chunks else [text[:max_len]]


def embed_texts(texts):
    """OpenAI embedding авах"""
    client = openai_client
    if not client:
        return None
    resp = client.embeddings.create(model="text-embedding-3-small", input=texts)
    return [d.embedding for d in resp.data]


@app.route("/api/embedding/store", methods=["POST"])
@editor_required
def store_embedding():
    """Тендерийн OCR текстийг embedding болгож хадгалах"""
    data = request.get_json()
    tender_no = data.get("tender_no", "").strip()

    if not tender_no:
        return jsonify({"error": "tender_no шаардлагатай"}), 400

    # tender_details-ээс текст авах
    detail = query("SELECT requirements, ocr_text, fail_reason FROM tender_details WHERE tender_no = %s",
                   (tender_no,), fetchone=True)
    if not detail:
        return jsonify({"error": "Тендерийн мэдээлэл олдсонгүй"}), 404

    text = (detail.get("requirements") or detail.get("ocr_text") or "").strip()
    if not text:
        return jsonify({"error": "Текст хоосон"}), 400

    chunks = chunk_text(text)
    embeddings = embed_texts(chunks)
    if embeddings is None:
        return jsonify({"error": "OPENAI_API_KEY тохируулаагүй"}), 500

    # Хуучныг устгаж шинээр хадгалах
    conn = get_db()
    cur = conn.cursor()
    cur.execute("DELETE FROM tender_embeddings WHERE tender_no = %s", (tender_no,))
    for i, (chunk, emb) in enumerate(zip(chunks, embeddings)):
        cur.execute("""
            INSERT INTO tender_embeddings (tender_no, chunk_index, chunk_text, embedding, source)
            VALUES (%s, %s, %s, %s, 'ocr')
        """, (tender_no, i, chunk, str(emb)))
    conn.commit()
    release_db(conn)

    return jsonify({"success": True, "chunks": len(chunks), "tender_no": tender_no})


@app.route("/api/embedding/search", methods=["POST"])
@login_required
def search_embedding():
    """Семантик хайлт — асуултыг embedding болгож ойролцоо тендер хайх"""
    data = request.get_json()
    query_text = data.get("query", "").strip()
    limit = int(data.get("limit", 10))

    if not query_text:
        return jsonify({"error": "Хайлтын текст оруулна уу"}), 400

    q_emb = embed_texts([query_text])
    if q_emb is None:
        return jsonify({"error": "OPENAI_API_KEY тохируулаагүй"}), 500

    rows = query("""
        SELECT tender_no, chunk_text, chunk_index,
               1 - (embedding <=> %s::vector) as similarity
        FROM tender_embeddings
        ORDER BY embedding <=> %s::vector
        LIMIT %s
    """, (str(q_emb[0]), str(q_emb[0]), limit), fetchall=True)

    from decimal import Decimal
    for r in rows:
        for k, v in r.items():
            if isinstance(v, Decimal):
                r[k] = float(v)

    return jsonify({"results": rows})


@app.route("/api/embedding/batch", methods=["POST"])
@editor_required
def batch_embedding():
    """Бүх tender_details-ийн текстийг embedding болгох"""
    details = query("""
        SELECT tender_no, COALESCE(requirements, '') || ' ' || COALESCE(ocr_text, '') as text
        FROM tender_details
        WHERE tender_no NOT IN (SELECT DISTINCT tender_no FROM tender_embeddings)
        LIMIT 50
    """, fetchall=True)

    if not details:
        return jsonify({"message": "Бүх текст embedding хийгдсэн", "processed": 0})

    processed = 0
    errors = 0
    for d in details:
        text = d["text"].strip()
        if not text or len(text) < 10:
            continue
        try:
            chunks = chunk_text(text)
            embeddings = embed_texts(chunks)
            if not embeddings:
                continue

            conn = get_db()
            cur = conn.cursor()
            for i, (chunk, emb) in enumerate(zip(chunks, embeddings)):
                cur.execute("""
                    INSERT INTO tender_embeddings (tender_no, chunk_index, chunk_text, embedding, source)
                    VALUES (%s, %s, %s, %s, 'ocr')
                """, (d["tender_no"], i, chunk, str(emb)))
            conn.commit()
            release_db(conn)
            processed += 1
        except Exception as e:
            errors += 1

    return jsonify({"processed": processed, "errors": errors, "remaining": len(details) - processed})


@app.route("/api/ai/cross-analysis", methods=["POST"])
@login_required
def cross_analysis():
    """В урсгал: Шаардлагын заалт × pgvector chunk → Claude cross-analysis
    Тендер бүрт шаардлага хангасан/хангаагүй дүгнэлт гаргах"""
    client = claude_client
    if not client:
        return jsonify({"error": "ANTHROPIC_API_KEY тохируулаагүй"}), 500

    data = request.get_json()
    tender_no = data.get("tender_no", "").strip()

    if not tender_no:
        return jsonify({"error": "tender_no шаардлагатай"}), 400

    # 1) Шаардлагын заалтууд авах
    detail = query("SELECT requirements, fail_reason FROM tender_details WHERE tender_no = %s",
                   (tender_no,), fetchone=True)
    if not detail or not detail.get("requirements"):
        return jsonify({"error": "Шаардлага оруулаагүй байна. Эхлээд PDF оруулна уу."}), 400

    requirements = detail["requirements"]
    fail_reason = detail.get("fail_reason") or ""

    # 2) pgvector-ээс холбогдох chunk-ууд авах
    chunks = query("""
        SELECT chunk_text, chunk_index FROM tender_embeddings
        WHERE tender_no = %s ORDER BY chunk_index
    """, (tender_no,), fetchall=True)

    chunk_texts = "\n".join([c["chunk_text"] for c in chunks]) if chunks else requirements

    # 3) Тендерийн мэдээлэл
    tender = query("SELECT * FROM joined_tenders WHERE tender_no = %s", (tender_no,), fetchone=True)
    tender_info = ""
    if tender:
        tender_info = f"Тендер: {tender.get('name','')} | Байгууллага: {tender.get('organization','')} | Статус: {tender.get('status','')}"

    # 4) Claude cross-analysis
    msg = client.messages.create(
        model=CLAUDE_MODEL,
        max_tokens=3000,
        messages=[{
            "role": "user",
            "content": f"""Тендерийн шаардлага болон унасан шалтгааг харьцуулж дүн шинжилгээ хий.

{tender_info}

=== ТЕНДЕРИЙН ШААРДЛАГУУД ===
{requirements[:3000]}

=== БАРИМТ БИЧГИЙН ДЭЛГЭРЭНГҮЙ (OCR) ===
{chunk_texts[:3000]}

=== УНАСАН ШАЛТГААН (хэрэв байгаа бол) ===
{fail_reason[:1000]}

ДААЛГАВАР:
Шаардлага бүрийг шинжилж, компани хангаж чадсан/чадаагүй эсэхийг дүгнэ.

JSON хариул:
{{
  "summary": "Ерөнхий дүгнэлт",
  "total_requirements": тоо,
  "met": хангасан_тоо,
  "not_met": хангаагүй_тоо,
  "unclear": тодорхойгүй_тоо,
  "analysis": [
    {{
      "id": "заалтын дугаар",
      "requirement": "шаардлагын товч текст",
      "status": "met/not_met/unclear",
      "reason": "Яагаад хангасан/хангаагүй тайлбар",
      "recommendation": "Цаашид юу хийх вэ"
    }}
  ],
  "key_lessons": ["Сургамж 1", "Сургамж 2"]
}}
Зөвхөн JSON хариул."""
        }]
    )

    raw = msg.content[0].text.strip()
    result = None
    try:
        result = json.loads(raw)
    except (json.JSONDecodeError, ValueError):
        start = raw.find('{')
        end = raw.rfind('}')
        if start >= 0 and end > start:
            try:
                result = json.loads(raw[start:end+1])
            except (json.JSONDecodeError, ValueError):
                pass
    if not result:
        result = {"summary": raw[:500], "analysis": []}

    # pgvector-т дүгнэлт хадгалах
    try:
        analysis_text = result.get("summary", "") + " " + " ".join([
            f"{a.get('id','')} {a.get('requirement','')} {a.get('status','')} {a.get('reason','')}"
            for a in result.get("analysis", [])
        ])
        if analysis_text.strip():
            embs = embed_texts([analysis_text[:2000]])
            if embs:
                conn = get_db()
                cur = conn.cursor()
                cur.execute("DELETE FROM tender_embeddings WHERE tender_no = %s AND source = 'analysis'", (tender_no,))
                cur.execute("""INSERT INTO tender_embeddings (tender_no, chunk_index, chunk_text, embedding, source)
                              VALUES (%s, 0, %s, %s, 'analysis')""",
                            (tender_no, analysis_text[:2000], str(embs[0])))
                conn.commit()
                release_db(conn)
    except Exception as e:
        logging.warning("Cross-analysis embedding алдаа: %s", e)

    result["tender_no"] = tender_no
    return jsonify(result)


@app.route("/api/embedding/stats")
@login_required
def embedding_stats():
    stats = query("""
        SELECT COUNT(DISTINCT tender_no) as tenders, COUNT(*) as chunks
        FROM tender_embeddings
    """, fetchone=True)
    return jsonify(stats)


# ============================================================
# BATCH OCR — Фолдероос автомат уншуулах
# ============================================================

@app.route("/api/ocr/batch", methods=["POST"])
@editor_required
def ocr_batch():
    """Фолдер доторх PDF файлуудыг тендер дугаараар таниж автомат OCR хийх.
    Файлын нэр: тендердугаар_шаардлага.pdf эсвэл тендердугаар_унасан.pdf
    """
    data = request.get_json()
    folder_path = data.get("folder", "").strip()

    if not folder_path or not os.path.isdir(folder_path):
        return jsonify({"error": f"Фолдер олдсонгүй: {folder_path}"}), 400

    files = [f for f in os.listdir(folder_path) if f.lower().endswith(('.pdf', '.png', '.jpg', '.jpeg'))]
    if not files:
        return jsonify({"error": "Фолдерт PDF/зураг файл олдсонгүй"}), 400

    client = claude_client
    user_id = session["user_id"]
    results = []
    processed = 0
    errors = 0

    for fname in sorted(files):
        # Нэрээс тендер дугаар + төрөл задлах
        name_part = os.path.splitext(fname)[0]  # .pdf хасах

        if '_шаардлага' in name_part:
            tender_no = name_part.replace('_шаардлага', '').strip()
            doc_type = 'requirements'
        elif '_унасан' in name_part:
            tender_no = name_part.replace('_унасан', '').strip()
            doc_type = 'fail_reason'
        else:
            results.append({"file": fname, "status": "skipped", "reason": "Нэр таниулгүй (_шаардлага эсвэл _унасан байх ёстой)"})
            continue

        # / тэмдэгт файлын нэрэнд байж болохгүй тул _ -аар солигдсон байж болно
        tender_no = tender_no.replace('_', '/', 1)

        # OCR
        filepath = os.path.join(folder_path, fname)
        try:
            raw_text = ""
            if fname.lower().endswith('.pdf'):
                with open(filepath, 'rb') as f:
                    images = convert_from_bytes(f.read(), dpi=300)
                for img in images:
                    raw_text += pytesseract.image_to_string(img, lang="mon+eng") + "\n"
            else:
                img = Image.open(filepath)
                raw_text = pytesseract.image_to_string(img, lang="mon+eng")

            if not raw_text.strip():
                results.append({"file": fname, "tender_no": tender_no, "type": doc_type, "status": "error", "reason": "Текст уншигдсангүй"})
                errors += 1
                continue

            # Claude засварлах (байвал)
            final_text = raw_text
            if client:
                try:
                    msg = client.messages.create(
                        model=CLAUDE_MODEL,
                        max_tokens=2000,
                        messages=[{
                            "role": "user",
                            "content": f"Дараах OCR текстийн Монгол кирилл алдааг засварла. Зөвхөн засварласан текстийг хариул:\n\n{raw_text[:3000]}"
                        }]
                    )
                    final_text = msg.content[0].text.strip()
                except Exception as e:
                    logging.warning("Batch OCR Claude засвар алдаа: %s", e)

            # DB хадгалах
            conn = get_db()
            cur = conn.cursor()
            cur.execute("SELECT id FROM tender_details WHERE tender_no = %s", (tender_no,))
            existing = cur.fetchone()

            if existing:
                if doc_type == 'requirements':
                    cur.execute("UPDATE tender_details SET requirements=%s, ocr_text=%s, updated_by=%s, updated_at=NOW() WHERE tender_no=%s",
                                (final_text, raw_text, user_id, tender_no))
                else:
                    cur.execute("UPDATE tender_details SET fail_reason=%s, updated_by=%s, updated_at=NOW() WHERE tender_no=%s",
                                (final_text, user_id, tender_no))
            else:
                if doc_type == 'requirements':
                    cur.execute("INSERT INTO tender_details (tender_no, requirements, ocr_text, updated_by) VALUES (%s,%s,%s,%s)",
                                (tender_no, final_text, raw_text, user_id))
                else:
                    cur.execute("INSERT INTO tender_details (tender_no, fail_reason, updated_by) VALUES (%s,%s,%s)",
                                (tender_no, final_text, user_id))

            conn.commit()
            release_db(conn)
            processed += 1
            results.append({"file": fname, "tender_no": tender_no, "type": doc_type, "status": "ok"})

        except Exception as e:
            errors += 1
            results.append({"file": fname, "tender_no": tender_no, "type": doc_type, "status": "error", "reason": str(e)})

    return jsonify({
        "success": True,
        "total_files": len(files),
        "processed": processed,
        "errors": errors,
        "skipped": len(files) - processed - errors,
        "results": results
    })


# ============================================================
# АДМИН — ӨГӨГДӨЛ ЦЭВЭРЛЭХ
# ============================================================

@app.route("/api/admin/clear/<table_name>", methods=["DELETE"])
@login_required
def clear_table(table_name):
    if session.get("role") != "admin":
        return jsonify({"error": "Зөвхөн админ"}), 403

    allowed = {
        "tenders": "tenders",
        "procurement_plans": "procurement_plans",
        "joined_tenders": "joined_tenders",
        "won_tenders": "won_tenders",
        "upload_history": "upload_history",
        "procurement_upload_history": "procurement_upload_history",
    }
    if table_name not in allowed:
        return jsonify({"error": "Зөвшөөрөгдөөгүй хүснэгт"}), 400

    conn = get_db()
    cur = conn.cursor()
    cur.execute(f"SELECT COUNT(*) FROM {allowed[table_name]}")
    count = cur.fetchone()[0]
    cur.execute(f"DELETE FROM {allowed[table_name]}")
    conn.commit()
    release_db(conn)

    return jsonify({"success": True, "deleted": count, "table": table_name})


@app.route("/api/admin/table-stats")
@login_required
def table_stats():
    if session.get("role") != "admin":
        return jsonify({"error": "Зөвхөн админ"}), 403

    tables = ["tenders", "procurement_plans", "joined_tenders", "won_tenders", "upload_history", "procurement_upload_history"]
    result = []
    for t in tables:
        row = query(f"SELECT COUNT(*) as cnt FROM {t}", fetchone=True)
        result.append({"table": t, "count": row["cnt"]})
    return jsonify(result)


@app.route("/api/admin/export/<table_name>")
@login_required
def export_table(table_name):
    if session.get("role") != "admin":
        return jsonify({"error": "Зөвхөн админ"}), 403

    allowed = ["tenders", "procurement_plans", "joined_tenders", "won_tenders", "upload_history", "procurement_upload_history"]
    if table_name not in allowed:
        return jsonify({"error": "Зөвшөөрөгдөөгүй"}), 400

    from decimal import Decimal

    conn = get_db()
    cur = conn.cursor(cursor_factory=psycopg2.extras.RealDictCursor)
    cur.execute(f"SELECT * FROM {table_name} ORDER BY id")
    rows = cur.fetchall()
    release_db(conn)

    if not rows:
        return jsonify({"error": "Хоосон хүснэгт"}), 404

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = table_name

    # Header
    headers = list(rows[0].keys())
    for col, h in enumerate(headers, 1):
        cell = ws.cell(row=1, column=col, value=h)
        cell.font = openpyxl.styles.Font(bold=True)

    # Data
    for r_idx, row in enumerate(rows, 2):
        for c_idx, h in enumerate(headers, 1):
            val = row[h]
            if isinstance(val, Decimal):
                val = float(val)
            elif val is not None and not isinstance(val, (int, float, datetime)):
                val = str(val)
            ws.cell(row=r_idx, column=c_idx, value=val)

    output = io.BytesIO()
    wb.save(output)
    output.seek(0)

    from flask import Response
    return Response(
        output.getvalue(),
        mimetype="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": f"attachment; filename={table_name}.xlsx"}
    )


# ============================================================
# AI — Claude Integration
# ============================================================



@app.route("/api/ai/match", methods=["POST"])
@login_required
def ai_match():
    """AI ашиглаж төлөвлөгөө <-> тендер нэр тулгах"""
    client = claude_client
    if not client:
        return jsonify({"error": "ANTHROPIC_API_KEY тохируулаагүй байна"}), 500

    data = request.get_json()
    plan_name = data.get("plan_name", "")
    candidates = data.get("candidates", [])  # [{tender_no, name}]

    if not plan_name or not candidates:
        return jsonify({"error": "plan_name болон candidates шаардлагатай"}), 400

    cand_text = "\n".join([f"- [{c['tender_no']}] {c['name']}" for c in candidates[:30]])

    msg = client.messages.create(
        model=CLAUDE_MODEL,
        max_tokens=500,
        messages=[{
            "role": "user",
            "content": f"""Дараах худалдан авалтын төлөвлөгөөний нэрийг тендерийн жагсаалтаас утга агуулгаар нь тааруулж өг.

Төлөвлөгөөний нэр: "{plan_name}"

Тендерийн жагсаалт:
{cand_text}

Хэрэв тааралт олдвол JSON-оор хариул: {{"match": true, "tender_no": "...", "confidence": 0.0-1.0, "reason": "..."}}
Хэрэв олдохгүй бол: {{"match": false, "reason": "..."}}
Зөвхөн JSON хариул, өөр юу ч бичих хэрэггүй."""
        }]
    )

    try:
        result = json.loads(msg.content[0].text)
    except (json.JSONDecodeError, ValueError):
        result = {"match": False, "reason": msg.content[0].text}

    return jsonify(result)


@app.route("/api/ai/analyze", methods=["POST"])
@login_required
def ai_analyze():
    """AI шинжилгээ — тендер/худалдан авалтын дата дээр асуулт асуух"""
    client = claude_client
    if not client:
        return jsonify({"error": "ANTHROPIC_API_KEY тохируулаагүй байна"}), 500

    data = request.get_json()
    question = data.get("question", "").strip()
    if not question:
        return jsonify({"error": "Асуулт оруулна уу"}), 400

    # Контекст бэлтгэх — тендер + төлөвлөгөөний статистик
    tender_stats = query("""
        SELECT COUNT(*) as total,
               COUNT(CASE WHEN deadline::timestamp >= NOW() THEN 1 END) as active,
               COUNT(CASE WHEN deadline::timestamp < NOW() THEN 1 END) as expired,
               COUNT(DISTINCT organization) as orgs
        FROM tenders
    """, fetchone=True)

    proc_stats = query("""
        SELECT COUNT(*) as total,
               COALESCE(SUM(budget_amount), 0) as total_budget,
               COUNT(DISTINCT organization) as orgs,
               COUNT(DISTINCT ts_type) as types
        FROM procurement_plans
    """, fetchone=True)

    # Сүүлийн тендерүүд
    recent = query("""
        SELECT tender_no, name, organization, type, method, deadline
        FROM tenders ORDER BY deadline DESC LIMIT 20
    """, fetchall=True)

    # Төрлөөр задаргаа
    type_breakdown = query("""
        SELECT type, COUNT(*) as cnt FROM tenders GROUP BY type
    """, fetchall=True)

    proc_types = query("""
        SELECT ts_type, COUNT(*) as cnt, COALESCE(SUM(budget_amount),0) as budget
        FROM procurement_plans GROUP BY ts_type
    """, fetchall=True)


    context = f"""Тендерийн мэдээлэл:
- Нийт тендер: {tender_stats['total']}, Идэвхтэй: {tender_stats['active']}, Хугацаа дууссан: {tender_stats['expired']}, Байгууллага: {tender_stats['orgs']}
- Төрлөөр: {json.dumps(type_breakdown, default=dec, ensure_ascii=False)}

Худалдан авалтын төлөвлөгөө:
- Нийт: {proc_stats['total']}, Нийт төсөв: {dec(proc_stats['total_budget'])} мян.₮, Байгууллага: {proc_stats['orgs']}
- Төрлөөр: {json.dumps(proc_types, default=dec, ensure_ascii=False)}

Сүүлийн 20 тендер:
{json.dumps(recent, default=dec, ensure_ascii=False)}"""

    msg = client.messages.create(
        model=CLAUDE_MODEL,
        max_tokens=1500,
        system="Чи тендер болон худалдан авалтын мэдээлэлд дүн шинжилгээ хийх AI туслах юм. Монгол хэлээр хариул. Товч, тодорхой, ашигтай мэдээлэл өг.",
        messages=[{
            "role": "user",
            "content": f"""Дараах өгөгдөл дээр суурилж асуултад хариул:

{context}

Асуулт: {question}"""
        }]
    )

    return jsonify({
        "answer": msg.content[0].text,
        "model": msg.model,
        "tokens": msg.usage.output_tokens
    })


@app.route("/api/ai/tender-search", methods=["POST"])
@login_required
def ai_tender_search():
    """AI тендер хайх, санал болгох, шүүх"""
    client = claude_client
    if not client:
        return jsonify({"error": "ANTHROPIC_API_KEY тохируулаагүй"}), 500

    data = request.get_json()
    question = data.get("question", "").strip()
    if not question:
        return jsonify({"error": "Асуулт оруулна уу"}), 400

    # Идэвхтэй тендерүүдээс авах
    tenders = query("""
        SELECT tender_no, name, organization, type, method, deadline, link
        FROM tenders WHERE deadline::timestamp >= NOW()
        ORDER BY deadline ASC LIMIT 300
    """, fetchall=True)

    tender_stats = query("""
        SELECT COUNT(*) as total,
               COUNT(CASE WHEN deadline::timestamp >= NOW() THEN 1 END) as active,
               COUNT(DISTINCT organization) as orgs,
               COUNT(DISTINCT type) as types
        FROM tenders
    """, fetchone=True)

    tender_list = "\n".join([
        f"[{t['tender_no']}] {t['name']} | {t['organization']} | {t['type']} | {t['method']} | дуусах: {t['deadline']}"
        for t in tenders[:200]
    ])

    msg = client.messages.create(
        model=CLAUDE_MODEL,
        max_tokens=1500,
        system="""Чи тендерийн мэдээллийн AI туслах юм. Монгол хэлээр хариул.
Хэрэглэгч тендер хайж, санал авч, мэдээлэл асууж болно.

Хэрэв хэрэглэгч тодорхой сэдвээр тендер хайвал тохирох тендерүүдийг JSON массиваар буцаа.
Хариултын төгсгөлд дараах форматаар холбогдох тендерүүдийг оруул:
###TENDERS###
[{"tender_no":"...","name":"...","organization":"...","type":"..."}]
###END###

Хэрэв ерөнхий асуулт бол зөвхөн текст хариул, TENDERS блок оруулах шаардлагагүй.""",
        messages=[{
            "role": "user",
            "content": f"""Тендерийн статистик: нийт {tender_stats['total']}, идэвхтэй {tender_stats['active']}, {tender_stats['orgs']} байгууллага

Идэвхтэй тендерүүд:
{tender_list}

Хэрэглэгчийн асуулт: {question}"""
        }]
    )

    answer_text = msg.content[0].text
    found_tenders = []

    # Parse TENDERS block
    if "###TENDERS###" in answer_text:
        parts = answer_text.split("###TENDERS###")
        answer_text = parts[0].strip()
        try:
            tender_json = parts[1].split("###END###")[0].strip()
            found_tenders = json.loads(tender_json)
        except (json.JSONDecodeError, ValueError, IndexError):
            pass

    return jsonify({
        "answer": answer_text,
        "tenders": found_tenders[:20],
        "model": msg.model
    })


@app.route("/api/profile/keywords")
@login_required
def get_profile_keywords():
    """Хадгалсан түлхүүр үгс + профайл мэдээлэл"""
    keywords = query("SELECT keyword, weight, category FROM profile_keywords ORDER BY weight DESC", fetchall=True)
    joined_count = query("SELECT COUNT(*) as cnt FROM joined_tenders", fetchone=True)["cnt"]
    won_count = query("SELECT COUNT(*) as cnt FROM won_tenders", fetchone=True)["cnt"]
    active_count = query("SELECT COUNT(*) as cnt FROM tenders WHERE deadline::timestamp >= NOW()", fetchone=True)["cnt"]

    from decimal import Decimal
    for k in keywords:
        if isinstance(k.get("weight"), Decimal):
            k["weight"] = float(k["weight"])

    return jsonify({
        "keywords": keywords,
        "joined_count": joined_count,
        "won_count": won_count,
        "active_count": active_count
    })


@app.route("/api/profile/generate", methods=["POST"])
@editor_required
def generate_profile():
    """А урсгал: Claude API-аар түлхүүр үг гаргаж profile_keywords-д хадгалах"""
    client = claude_client
    if not client:
        return jsonify({"error": "ANTHROPIC_API_KEY тохируулаагүй"}), 500

    joined = query("SELECT name, organization, tender_type FROM joined_tenders ORDER BY id DESC LIMIT 300", fetchall=True)
    won = query("SELECT name, tender_type FROM won_tenders LIMIT 100", fetchall=True)
    if not joined:
        return jsonify({"error": "Оролцсон тендер оруулаагүй"}), 400

    names = [j["name"] for j in joined if j["name"]]
    won_names = [w["name"] for w in won if w["name"]]

    msg = client.messages.create(
        model=CLAUDE_MODEL,
        max_tokens=2000,
        messages=[{
            "role": "user",
            "content": f"""Дараах тендерийн нэрсийг шинжилж компанийн чиглэлийн ТҮЛХҮҮР ҮГС гарга.

ОРОЛЦСОН ТЕНДЕРҮҮД:
{'; '.join(names[:100])}

ШАЛГАРСАН:
{'; '.join(won_names[:30])}

ДААЛГАВАР: Эдгээр тендерээс мэргэжлийн түлхүүр үгс гарга. Жишээ: "трансформатор", "реле хамгаалалт", "кабель", "засвар", "сэлбэг" гэх мэт.
Ерөнхий үгс (тендер, ажил, бараа, худалдан) ОРУУЛАХГҮЙ. Зөвхөн мэргэжлийн тодорхой үгс.

JSON хариул:
[{{"keyword":"трансформатор","weight":0.9,"category":"тоног төхөөрөмж"}},{{"keyword":"кабель","weight":0.8,"category":"материал"}}]
Зөвхөн JSON массив, 30-50 түлхүүр үг."""
        }]
    )

    raw = msg.content[0].text.strip()
    kw_list = None
    try:
        kw_list = json.loads(raw)
    except (json.JSONDecodeError, ValueError):
        start = raw.find('[')
        end = raw.rfind(']')
        if start >= 0 and end > start:
            try:
                kw_list = json.loads(raw[start:end+1])
            except (json.JSONDecodeError, ValueError):
                pass

    if not kw_list or not isinstance(kw_list, list):
        return jsonify({"error": "AI хариулт задлах чадсангүй"}), 500

    # DB-д хадгалах
    conn = get_db()
    cur = conn.cursor()
    cur.execute("DELETE FROM profile_keywords")
    inserted = 0
    for kw in kw_list:
        word = (kw.get("keyword") or "").strip().lower()
        if not word or len(word) < 2:
            continue
        weight = float(kw.get("weight", 1.0))
        cat = (kw.get("category") or "").strip()
        cur.execute("INSERT INTO profile_keywords (keyword, weight, category) VALUES (%s,%s,%s) ON CONFLICT (keyword) DO UPDATE SET weight=%s, category=%s",
                    (word, weight, cat, weight, cat))
        inserted += 1
    conn.commit()
    release_db(conn)

    return jsonify({"success": True, "count": inserted, "keywords": kw_list})


@app.route("/api/profile/keywords", methods=["PUT"])
@editor_required
def update_profile_keyword():
    """Түлхүүр үг нэмэх/засах/устгах"""
    data = request.get_json()
    action = data.get("action", "add")
    keyword = (data.get("keyword") or "").strip().lower()

    if not keyword:
        return jsonify({"error": "Түлхүүр үг оруулна уу"}), 400

    if action == "delete":
        query("DELETE FROM profile_keywords WHERE keyword = %s", (keyword,), commit=True)
        return jsonify({"success": True})
    else:
        weight = float(data.get("weight", 1.0))
        cat = (data.get("category") or "").strip()
        query("INSERT INTO profile_keywords (keyword, weight, category) VALUES (%s,%s,%s) ON CONFLICT (keyword) DO UPDATE SET weight=%s, category=%s",
              (keyword, weight, cat, weight, cat), commit=True)
        return jsonify({"success": True})


@app.route("/api/tenders/match")
@login_required
def match_tenders():
    """Б урсгал: profile_keywords ашиглаж SQL ILIKE-аар тендер шүүх. Claude дуудахгүй."""
    keywords = query("SELECT keyword, weight FROM profile_keywords ORDER BY weight DESC", fetchall=True)
    if not keywords:
        return jsonify({"error": "Профайл үүсгээгүй. Эхлээд 'Профайл үүсгэх' товч дарна уу."}), 400

    source = request.args.get("source", "tenders")  # tenders / plans
    limit = min(int(request.args.get("limit", 50)), 200)
    offset = int(request.args.get("offset", 0))

    from decimal import Decimal

    if source == "plans":
        rows = query("SELECT ts_code as id, name, organization, ts_type as type, budget_amount, procurement_method as method FROM procurement_plans", fetchall=True)
    else:
        rows = query("""
            SELECT tender_no as id, name, organization, type, method, deadline, link
            FROM tenders WHERE deadline::timestamp >= NOW()
            ORDER BY deadline ASC
        """, fetchall=True)

    # Keyword scoring
    results = []
    for r in rows:
        name_lower = (r.get("name") or "").lower()
        score = 0
        matched = []
        for kw in keywords:
            w = kw["keyword"]
            wt = float(kw["weight"]) if isinstance(kw["weight"], Decimal) else kw["weight"]
            if w in name_lower:
                score += wt
                matched.append(w)

        if score > 0:
            # Decimal -> float
            clean = {}
            for k, v in r.items():
                clean[k] = float(v) if isinstance(v, Decimal) else v

            clean["score"] = round(score, 2)
            clean["matched_keywords"] = matched
            clean["source"] = "PLAN" if source == "plans" else "TENDER"
            results.append(clean)

    # Score-оор эрэмбэлэх
    results.sort(key=lambda x: x["score"], reverse=True)

    # Normalize score to 0-100
    max_score = results[0]["score"] if results else 1
    for r in results:
        r["score_pct"] = min(99, round(r["score"] / max_score * 100))

    total = len(results)
    page_results = results[offset:offset+limit]

    return jsonify({"total": total, "results": page_results})


@app.route("/api/ai/bulk-match", methods=["POST"])
@editor_required
def ai_bulk_match():
    """Бүх төлөвлөгөөг тендертэй AI-аар тулгах"""
    client = claude_client
    if not client:
        return jsonify({"error": "ANTHROPIC_API_KEY тохируулаагүй байна"}), 500

    # Тулгагдаагүй төлөвлөгөөнүүдийг авах
    unmatched = query("""
        SELECT p.id, p.ts_code, p.name
        FROM procurement_plans p
        LEFT JOIN LATERAL (
            SELECT t.tender_no FROM tenders t
            WHERE p.ts_code = SPLIT_PART(t.tender_no, '/', 1) || '/' || SPLIT_PART(t.tender_no, '/', 2)
               OR LOWER(TRIM(p.name)) = LOWER(TRIM(t.name))
            LIMIT 1
        ) t ON true
        WHERE t.tender_no IS NULL
        LIMIT 50
    """, fetchall=True)

    if not unmatched:
        return jsonify({"message": "Бүх төлөвлөгөө тулгагдсан байна", "matched": 0})

    # Тендерүүдийн нэрсийг авах
    tenders = query("""
        SELECT tender_no, name FROM tenders
        WHERE deadline::timestamp >= NOW()
        ORDER BY deadline ASC LIMIT 500
    """, fetchall=True)

    tender_list = "\n".join([f"[{t['tender_no']}] {t['name']}" for t in tenders[:200]])

    results = []
    for plan in unmatched[:20]:  # 20-оор хязгаарлах (API cost)
        try:
            msg = client.messages.create(
                model=CLAUDE_MODEL,
                max_tokens=300,
                messages=[{
                    "role": "user",
                    "content": f"""Дараах худалдан авалтыг тендерийн жагсаалтаас утга, агуулгаар тааруул.

Худалдан авалт: "{plan['name']}" (код: {plan['ts_code']})

Тендерүүд:
{tender_list}

JSON хариул: {{"match": true/false, "tender_no": "...", "confidence": 0.0-1.0}} Зөвхөн JSON."""
                }]
            )
            r = json.loads(msg.content[0].text)
            r["plan_id"] = plan["id"]
            r["plan_name"] = plan["name"]
            results.append(r)
        except Exception as e:
            logging.warning("Bulk match алдаа [%s]: %s", plan.get("name", "?"), e)
            continue

    return jsonify({
        "total_unmatched": len(unmatched),
        "processed": len(results),
        "matches": [r for r in results if r.get("match")],
        "no_match": [r for r in results if not r.get("match")]
    })


# ============================================================
# DASHBOARD PAGE
# ============================================================

@app.route("/api/dashboard")
@login_required
def dashboard_overview():
    """Ерөнхий dashboard мэдээлэл"""
    from decimal import Decimal
    def dec(v):
        return float(v) if isinstance(v, Decimal) else v

    # Тендер
    t = query("SELECT COUNT(*) as total, COUNT(CASE WHEN deadline::timestamp >= NOW() THEN 1 END) as active, COUNT(DISTINCT organization) as orgs FROM tenders", fetchone=True)

    # Оролцсон
    j = query("""
        SELECT COUNT(*) as total, COUNT(DISTINCT organization) as orgs,
               COUNT(CASE WHEN j.status='Үр дүн гарсан' THEN 1 END) as completed
        FROM joined_tenders j
    """, fetchone=True)

    # Шалгарсан
    w = query("SELECT COUNT(*) as total FROM won_tenders", fetchone=True)

    # Төлөвлөгөө
    p = query("SELECT COUNT(*) as total, COALESCE(SUM(budget_amount),0) as budget, COUNT(DISTINCT organization) as orgs FROM procurement_plans", fetchone=True)

    # Топ 5 байгууллага (хамгийн олон шалгарсан)
    top_won = query("""
        SELECT w.organization, COUNT(*) as cnt
        FROM won_tenders w
        WHERE w.organization IS NOT NULL AND w.organization != ''
        GROUP BY w.organization ORDER BY cnt DESC LIMIT 5
    """, fetchall=True)

    # Топ 5 байгууллага (хамгийн олон оролцсон)
    top_joined = query("""
        SELECT organization, COUNT(*) as cnt
        FROM joined_tenders
        WHERE organization IS NOT NULL AND organization != ''
        GROUP BY organization ORDER BY cnt DESC LIMIT 5
    """, fetchall=True)

    # Тендер зарласан топ байгууллагууд
    top_tender_orgs = query("""
        SELECT organization, COUNT(*) as cnt
        FROM tenders
        WHERE organization IS NOT NULL AND organization != ''
        GROUP BY organization ORDER BY cnt DESC LIMIT 10
    """, fetchall=True)

    # Төрлөөр
    type_stats = query("SELECT type, COUNT(*) as cnt FROM tenders WHERE type IS NOT NULL GROUP BY type ORDER BY cnt DESC", fetchall=True)

    # Оролцсон төрлөөр
    joined_types = query("SELECT tender_type, COUNT(*) as cnt FROM joined_tenders WHERE tender_type IS NOT NULL GROUP BY tender_type ORDER BY cnt DESC", fetchall=True)

    # Ялалтын хувь
    win_rate = 0
    if j["completed"] and j["completed"] > 0:
        win_rate = round((w["total"] / j["completed"]) * 100, 1)

    # Embedding stats
    emb = query("SELECT COUNT(DISTINCT tender_no) as tenders, COUNT(*) as chunks FROM tender_embeddings", fetchone=True)

    # Сүүлийн үйл ажиллагаа
    recent_joined = query("SELECT tender_no, name, organization, status FROM joined_tenders ORDER BY id DESC LIMIT 5", fetchall=True)

    return jsonify({
        "tenders": {"total": t["total"], "active": t["active"], "orgs": t["orgs"]},
        "joined": {"total": j["total"], "orgs": j["orgs"], "completed": j["completed"]},
        "won": {"total": w["total"], "win_rate": win_rate},
        "procurement": {"total": p["total"], "budget": dec(p["budget"]), "orgs": p["orgs"]},
        "top_won_orgs": top_won,
        "top_joined_orgs": top_joined,
        "top_tender_orgs": top_tender_orgs,
        "type_stats": type_stats,
        "joined_types": joined_types,
        "embeddings": emb,
        "recent_joined": recent_joined
    })


@app.route("/")
def index():
    return send_file(OVERVIEW_PATH)

@app.route("/tenders")
def tenders_page():
    return send_file(DASHBOARD_PATH)

@app.route("/procurement")
def procurement_page():
    return send_file(PROCUREMENT_PATH)

@app.route("/joined")
def joined_page():
    return send_file(JOINED_PATH)

@app.route("/suggested")
def suggested_page():
    return send_file(SUGGESTED_PATH)


# ============================================================
# MAIN
# ============================================================

if __name__ == "__main__":
    init_db()
    port = int(os.environ.get("PORT", 5000))
    print(f"[tender] PostgreSQL: {DB_CONFIG['host']}:{DB_CONFIG['port']}/{DB_CONFIG['dbname']}")
    print(f"[tender] Сервер эхэллээ: http://localhost:{port}")
    print(f"[tender] Анхдагч нэвтрэх: admin / admin123")
    print(f"[tender] Эрхүүд: admin (бүх эрх), editor (мэдээлэл оруулах), viewer (зөвхөн харах)")
    print(f"[tender] Худалдан авалтын төлөвлөгөө: Excel upload дэмжигдэнэ")
    debug = os.environ.get("FLASK_DEBUG", "false").lower() == "true"
    app.run(host="0.0.0.0", port=port, debug=debug)
