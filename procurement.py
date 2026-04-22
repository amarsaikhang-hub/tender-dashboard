import io
import logging
from datetime import datetime
from decimal import Decimal

import openpyxl
import psycopg2.extras
from flask import request, jsonify, session

from db import get_db, release_db, query
from auth import login_required, editor_required


def register(app):
    @app.route("/api/procurement/upload", methods=["POST"])
    @editor_required
    def upload_procurement():
        if "file" not in request.files:
            return jsonify({"error": "Excel файл сонгоно уу"}), 400

        file = request.files["file"]
        if not file.filename.lower().endswith((".xlsx", ".xls")):
            return jsonify({"error": "Зөвхөн .xlsx файл дэмжигдэнэ"}), 400

        organization = request.form.get("organization", "").strip()
        plan_year    = int(request.form.get("plan_year", datetime.now().year))

        if not organization:
            return jsonify({"error": "Байгууллагын нэр оруулна уу"}), 400

        try:
            wb = openpyxl.load_workbook(io.BytesIO(file.read()), read_only=True, data_only=True)
            ws = wb["sheet1"] if "sheet1" in wb.sheetnames else wb[wb.sheetnames[0]]

            headers = [str(cell.value or "").strip() for cell in next(ws.iter_rows(min_row=1, max_row=1))]
            expected = {
                "№": "row_num", "ТШ код": "ts_code", "нэр төрөл": "name",
                "ТШ төрөл": "ts_type", "эх үүсвэр": "funding_source",
                "Төсөвт өртөг": "budget_amount", "журам": "procurement_method",
                "Тендер зарлах": "tender_announce_date", "эрх олгох": "contract_award_date",
                "дуусгавар": "contract_end_date", "санхүүжих дүн": "annual_funding",
                "Тогтвортой": "sustainable_criteria", "Тайлбар": "description",
                "Үүсгэсэн огноо": "created_date",
            }
            col_map = {}
            for i, h in enumerate(headers):
                for key, field in expected.items():
                    if key.lower() in h.lower():
                        col_map[field] = i
                        break

            sub_cats = {}
            for sname in ["Бараа", "Ажил", "Үйлчилгээ"]:
                if sname in wb.sheetnames:
                    sub_ws = wb[sname]
                    sub_headers = [str(c.value or "").strip()
                                   for c in next(sub_ws.iter_rows(min_row=1, max_row=1))]
                    code_idx = next((i for i, h in enumerate(sub_headers) if "код" in h.lower()), None)
                    sub_idx  = next((i for i, h in enumerate(sub_headers) if "ангилал" in h.lower()), None)
                    if code_idx is not None and sub_idx is not None:
                        for row in sub_ws.iter_rows(min_row=2, values_only=True):
                            code = str(row[code_idx] or "").strip()
                            sub  = str(row[sub_idx]  or "").strip()
                            if code and sub:
                                sub_cats[code] = sub

            conn    = get_db()
            cur     = conn.cursor()
            user_id = session["user_id"]
            now     = datetime.utcnow().isoformat()
            inserted = updated = total = 0

            for row in ws.iter_rows(min_row=2, values_only=True):
                vals = list(row)
                if not vals or len(vals) < 3:
                    continue

                ts_code = str(vals[col_map.get("ts_code", 1)] or "").strip()
                if not ts_code:
                    continue

                total += 1
                name       = str(vals[col_map.get("name",       2)]  or "").strip()
                ts_type    = str(vals[col_map.get("ts_type",    3)]  or "").strip()
                funding    = str(vals[col_map.get("funding_source", 4)] or "").strip()
                budget     = vals[col_map.get("budget_amount", 5)] or 0
                method     = str(vals[col_map.get("procurement_method", 6)] or "").strip()
                t_announce = str(vals[col_map.get("tender_announce_date", 7)] or "").strip()
                c_award    = str(vals[col_map.get("contract_award_date",  8)] or "").strip()
                c_end      = str(vals[col_map.get("contract_end_date",    9)] or "").strip()
                ann_fund   = vals[col_map.get("annual_funding", 10)] or 0
                sustain    = str(vals[col_map.get("sustainable_criteria", 11)] or "").strip()
                desc       = str(vals[col_map.get("description",  12)] or "").strip()
                c_date     = str(vals[col_map.get("created_date", 13)] or "").strip()
                row_num_val = vals[col_map.get("row_num", 0)] or 0

                try:    budget      = float(budget)
                except: budget      = 0
                try:    ann_fund    = float(ann_fund)
                except: ann_fund    = 0
                try:    row_num_val = int(row_num_val)
                except: row_num_val = 0

                sub_cat = sub_cats.get(ts_code, "")

                cur.execute(
                    "SELECT id FROM procurement_plans WHERE ts_code=%s AND organization=%s AND plan_year=%s",
                    (ts_code, organization, plan_year)
                )
                if cur.fetchone():
                    cur.execute("""
                        UPDATE procurement_plans SET
                            row_num=%s, name=%s, ts_type=%s, funding_source=%s, budget_amount=%s,
                            procurement_method=%s, tender_announce_date=%s, contract_award_date=%s,
                            contract_end_date=%s, annual_funding=%s, sustainable_criteria=%s,
                            description=%s, created_date=%s, sub_category=%s, updated_at=%s
                        WHERE ts_code=%s AND organization=%s AND plan_year=%s
                    """, (row_num_val, name, ts_type, funding, budget, method, t_announce,
                          c_award, c_end, ann_fund, sustain, desc, c_date, sub_cat, now,
                          ts_code, organization, plan_year))
                    updated += 1
                else:
                    cur.execute("""
                        INSERT INTO procurement_plans
                            (row_num, ts_code, name, ts_type, funding_source, budget_amount,
                             procurement_method, tender_announce_date, contract_award_date,
                             contract_end_date, annual_funding, sustainable_criteria,
                             description, created_date, sub_category, organization,
                             plan_year, created_by, created_at, updated_at)
                        VALUES (%s,%s,%s,%s,%s,%s,%s,%s,%s,%s,%s,%s,%s,%s,%s,%s,%s,%s,%s,%s)
                    """, (row_num_val, ts_code, name, ts_type, funding, budget, method,
                          t_announce, c_award, c_end, ann_fund, sustain, desc, c_date,
                          sub_cat, organization, plan_year, user_id, now, now))
                    inserted += 1

            cur.execute("""
                INSERT INTO procurement_upload_history
                    (user_id, filename, organization, plan_year, total_rows, inserted, updated)
                VALUES (%s,%s,%s,%s,%s,%s,%s)
            """, (user_id, file.filename, organization, plan_year, total, inserted, updated))
            conn.commit()
            release_db(conn)
            wb.close()

            return jsonify({
                "success": True,
                "summary": {"total": total, "inserted": inserted, "updated": updated,
                            "organization": organization, "plan_year": plan_year}
            })
        except Exception as e:
            return jsonify({"error": f"Файл боловсруулж чадсангүй: {str(e)}"}), 500

    @app.route("/api/procurement/pending")
    @login_required
    def procurement_pending():
        search = request.args.get("search", "")
        ts_type = request.args.get("type", "")
        org     = request.args.get("organization", "")
        year    = request.args.get("year", "")
        status  = request.args.get("status", "")
        limit   = min(int(request.args.get("limit", 100)), 5000)
        offset  = int(request.args.get("offset", 0))

        conn = get_db()
        cur  = conn.cursor(cursor_factory=psycopg2.extras.RealDictCursor)

        try:
            cur.execute("""
                CREATE TEMP TABLE matched_plans ON COMMIT DROP AS
                SELECT p.*,
                       t.tender_no as matched_tender, t.deadline as tender_deadline,
                       t.method as tender_method, t.link as tender_link,
                       CASE WHEN t.tender_no IS NULL THEN 'not_announced' ELSE 'active' END as tender_status
                FROM procurement_plans p
                LEFT JOIN (
                    SELECT DISTINCT ON (code)
                        SPLIT_PART(tender_no, '/', 1) || '/' || SPLIT_PART(tender_no, '/', 2) as code,
                        tender_no, deadline, method, link, name
                    FROM tenders ORDER BY code, deadline DESC
                ) t ON (p.ts_code = t.code OR LOWER(TRIM(p.name)) = LOWER(TRIM(t.name)))
                WHERE (t.tender_no IS NULL OR t.deadline::timestamp >= NOW())
            """)

            where, params = ["1=1"], []
            if search:
                where.append("(name ILIKE %s OR ts_code ILIKE %s)")
                params.extend([f"%{search}%"] * 2)
            if ts_type:
                where.append("ts_type = %s"); params.append(ts_type)
            if org:
                where.append("organization = %s"); params.append(org)
            if year:
                where.append("plan_year = %s"); params.append(int(year))
            if status:
                where.append("tender_status = %s"); params.append(status)

            where_sql = " AND ".join(where)

            cur.execute(f"""
                SELECT COUNT(*) as total,
                       COUNT(CASE WHEN tender_status='not_announced' THEN 1 END) as not_announced,
                       COUNT(CASE WHEN tender_status='active' THEN 1 END) as announced
                FROM matched_plans WHERE {where_sql}
            """, params)
            stats = cur.fetchone()

            cur.execute(f"""
                SELECT * FROM matched_plans WHERE {where_sql}
                ORDER BY row_num ASC LIMIT %s OFFSET %s
            """, params + [limit, offset])
            rows = cur.fetchall()

            for r in rows:
                for k, v in r.items():
                    if isinstance(v, Decimal):
                        r[k] = float(v)

            conn.commit()
        except Exception:
            conn.rollback()
            raise
        finally:
            release_db(conn)

        return jsonify({
            "total": stats["total"], "plans": rows,
            "stats": {"total": stats["total"], "not_announced": stats["not_announced"],
                      "announced": stats["announced"]}
        })

    @app.route("/api/procurement")
    @login_required
    def get_procurement():
        org    = request.args.get("organization", "")
        year   = request.args.get("year", "")
        ts_type = request.args.get("type", "")
        search = request.args.get("search", "")
        limit  = min(int(request.args.get("limit", 5000)), 10000)
        offset = int(request.args.get("offset", 0))

        where, params = [], []
        if org:    where.append("organization = %s");  params.append(org)
        if year:   where.append("plan_year = %s");     params.append(int(year))
        if ts_type: where.append("ts_type = %s");      params.append(ts_type)
        if search:
            where.append("(name ILIKE %s OR ts_code ILIKE %s OR sub_category ILIKE %s)")
            params.extend([f"%{search}%"] * 3)

        where_sql = " AND ".join(where) if where else "1=1"
        total = query(f"SELECT COUNT(*) as cnt FROM procurement_plans WHERE {where_sql}",
                      params, fetchone=True)["cnt"]
        rows  = query(f"""
            SELECT * FROM procurement_plans WHERE {where_sql}
            ORDER BY row_num ASC LIMIT %s OFFSET %s
        """, params + [limit, offset], fetchall=True)

        for r in rows:
            for k, v in r.items():
                if isinstance(v, Decimal):
                    r[k] = float(v)

        return jsonify({"total": total, "plans": rows})

    @app.route("/api/procurement/stats")
    @login_required
    def procurement_stats():
        orgs  = query("SELECT DISTINCT organization FROM procurement_plans ORDER BY organization", fetchall=True)
        years = query("SELECT DISTINCT plan_year FROM procurement_plans ORDER BY plan_year DESC", fetchall=True)
        types = query("""
            SELECT ts_type, COUNT(*) as cnt, COALESCE(SUM(budget_amount),0) as total_budget
            FROM procurement_plans GROUP BY ts_type ORDER BY cnt DESC
        """, fetchall=True)
        total = query(
            "SELECT COUNT(*) as cnt, COALESCE(SUM(budget_amount),0) as total_budget FROM procurement_plans",
            fetchone=True
        )

        for t in types:
            for k, v in t.items():
                if isinstance(v, Decimal): t[k] = float(v)
        for k, v in total.items():
            if isinstance(v, Decimal): total[k] = float(v)

        return jsonify({
            "organizations": [r["organization"] for r in orgs],
            "years":         [r["plan_year"]    for r in years],
            "types":         types,
            "total_count":   total["cnt"],
            "total_budget":  total["total_budget"]
        })

    @app.route("/api/procurement/history")
    @login_required
    def procurement_history():
        rows = query("""
            SELECT h.*, u.full_name FROM procurement_upload_history h
            JOIN users u ON h.user_id = u.id
            ORDER BY h.uploaded_at DESC LIMIT 50
        """, fetchall=True)
        return jsonify(rows)
