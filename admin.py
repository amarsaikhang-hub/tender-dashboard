import io
from datetime import datetime
from decimal import Decimal

import openpyxl
import psycopg2.extras
from flask import request, jsonify, session, Response, send_file

from db import get_db, release_db, query
from auth import login_required
from config import dec, OVERVIEW_PATH, DASHBOARD_PATH, PROCUREMENT_PATH, JOINED_PATH, SUGGESTED_PATH

_ALLOWED_TABLES = {
    "tenders":                    "tenders",
    "procurement_plans":          "procurement_plans",
    "joined_tenders":             "joined_tenders",
    "won_tenders":                "won_tenders",
    "upload_history":             "upload_history",
    "procurement_upload_history": "procurement_upload_history",
}


def register(app):
    @app.route("/api/admin/clear/<table_name>", methods=["DELETE"])
    @login_required
    def clear_table(table_name):
        if session.get("role") != "admin":
            return jsonify({"error": "Зөвхөн админ"}), 403
        if table_name not in _ALLOWED_TABLES:
            return jsonify({"error": "Зөвшөөрөгдөөгүй хүснэгт"}), 400

        conn = get_db()
        cur  = conn.cursor()
        cur.execute(f"SELECT COUNT(*) FROM {_ALLOWED_TABLES[table_name]}")
        count = cur.fetchone()[0]
        cur.execute(f"DELETE FROM {_ALLOWED_TABLES[table_name]}")
        conn.commit()
        release_db(conn)
        return jsonify({"success": True, "deleted": count, "table": table_name})

    @app.route("/api/admin/table-stats")
    @login_required
    def table_stats():
        if session.get("role") != "admin":
            return jsonify({"error": "Зөвхөн админ"}), 403
        result = [{"table": t, "count": query(f"SELECT COUNT(*) as cnt FROM {t}", fetchone=True)["cnt"]}
                  for t in _ALLOWED_TABLES]
        return jsonify(result)

    @app.route("/api/admin/export/<table_name>")
    @login_required
    def export_table(table_name):
        if session.get("role") != "admin":
            return jsonify({"error": "Зөвхөн админ"}), 403
        if table_name not in _ALLOWED_TABLES:
            return jsonify({"error": "Зөвшөөрөгдөөгүй"}), 400

        conn = get_db()
        cur  = conn.cursor(cursor_factory=psycopg2.extras.RealDictCursor)
        cur.execute(f"SELECT * FROM {table_name} ORDER BY id")
        rows = cur.fetchall()
        release_db(conn)

        if not rows:
            return jsonify({"error": "Хоосон хүснэгт"}), 404

        wb      = openpyxl.Workbook()
        ws      = wb.active
        ws.title = table_name
        headers  = list(rows[0].keys())

        for col, h in enumerate(headers, 1):
            ws.cell(row=1, column=col, value=h).font = openpyxl.styles.Font(bold=True)

        for r_idx, row in enumerate(rows, 2):
            for c_idx, h in enumerate(headers, 1):
                val = row[h]
                if isinstance(val, Decimal): val = float(val)
                elif val is not None and not isinstance(val, (int, float, datetime)): val = str(val)
                ws.cell(row=r_idx, column=c_idx, value=val)

        output = io.BytesIO()
        wb.save(output)
        output.seek(0)
        return Response(
            output.getvalue(),
            mimetype="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
            headers={"Content-Disposition": f"attachment; filename={table_name}.xlsx"}
        )

    @app.route("/api/dashboard")
    @login_required
    def dashboard_overview():
        t = query("""
            SELECT COUNT(*) as total,
                   COUNT(CASE WHEN deadline::timestamp >= NOW() THEN 1 END) as active,
                   COUNT(DISTINCT organization) as orgs
            FROM tenders
        """, fetchone=True)
        j = query("""
            SELECT COUNT(*) as total, COUNT(DISTINCT organization) as orgs,
                   COUNT(CASE WHEN j.status='Үр дүн гарсан' THEN 1 END) as completed
            FROM joined_tenders j
        """, fetchone=True)
        w = query("SELECT COUNT(*) as total FROM won_tenders", fetchone=True)
        p = query("""
            SELECT COUNT(*) as total, COALESCE(SUM(budget_amount),0) as budget,
                   COUNT(DISTINCT organization) as orgs
            FROM procurement_plans
        """, fetchone=True)

        top_won        = query("""
            SELECT organization, COUNT(*) as cnt FROM won_tenders
            WHERE organization IS NOT NULL AND organization != ''
            GROUP BY organization ORDER BY cnt DESC LIMIT 5
        """, fetchall=True)
        top_joined     = query("""
            SELECT organization, COUNT(*) as cnt FROM joined_tenders
            WHERE organization IS NOT NULL AND organization != ''
            GROUP BY organization ORDER BY cnt DESC LIMIT 5
        """, fetchall=True)
        top_tender_orgs = query("""
            SELECT organization, COUNT(*) as cnt FROM tenders
            WHERE organization IS NOT NULL AND organization != ''
            GROUP BY organization ORDER BY cnt DESC LIMIT 10
        """, fetchall=True)
        type_stats     = query("""
            SELECT type, COUNT(*) as cnt FROM tenders
            WHERE type IS NOT NULL GROUP BY type ORDER BY cnt DESC
        """, fetchall=True)
        joined_types   = query("""
            SELECT tender_type, COUNT(*) as cnt FROM joined_tenders
            WHERE tender_type IS NOT NULL GROUP BY tender_type ORDER BY cnt DESC
        """, fetchall=True)
        emb            = query("""
            SELECT COUNT(DISTINCT tender_no) as tenders, COUNT(*) as chunks FROM tender_embeddings
        """, fetchone=True)
        recent_joined  = query("""
            SELECT tender_no, name, organization, status FROM joined_tenders ORDER BY id DESC LIMIT 5
        """, fetchall=True)

        win_rate = round((w["total"] / j["completed"]) * 100, 1) if j.get("completed") else 0

        return jsonify({
            "tenders":         {"total": t["total"], "active": t["active"], "orgs": t["orgs"]},
            "joined":          {"total": j["total"], "orgs": j["orgs"], "completed": j["completed"]},
            "won":             {"total": w["total"], "win_rate": win_rate},
            "procurement":     {"total": p["total"], "budget": dec(p["budget"]), "orgs": p["orgs"]},
            "top_won_orgs":    top_won,
            "top_joined_orgs": top_joined,
            "top_tender_orgs": top_tender_orgs,
            "type_stats":      type_stats,
            "joined_types":    joined_types,
            "embeddings":      emb,
            "recent_joined":   recent_joined
        })

    # ── HTML хуудаснууд ──────────────────────────────────────────
    @app.route("/")
    def index():
        return send_file(OVERVIEW_PATH)

    @app.route("/tenders")
    def tenders_page():
        return send_file(DASHBOARD_PATH)

    @app.route("/procurement")
    def procurement_page():
        return send_file(PROCUREMENT_PATH)

    @app.route("/joined")
    def joined_page():
        return send_file(JOINED_PATH)

    @app.route("/suggested")
    def suggested_page():
        return send_file(SUGGESTED_PATH)
